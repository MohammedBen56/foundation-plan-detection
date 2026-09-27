#!/usr/bin/env python3
"""Geometry-first classification of the elements on a foundation plan.

1. objects   every drawn object from geometry alone: solid fills (grouped by paint order),
             cross-hatched regions, closed outlines, bands (pairs of parallel lines, widths found
             from the drawing itself), stamp circles with corner rays. No colour, layer or text.
2. cluster   objects of the same drawing kind are clustered on geometric features (size,
             shape, what they contain, what they touch). Still no text.
3. match     labels are matched to objects by distance and text alignment in one global
             assignment, whatever their type; a label may stay unmatched.
4. name      each cluster takes the name of the label type most of its matched labels carry;
             unlabelled members inherit it; labels that disagree are listed as conflicts.
5. re-match  labels are matched again, now only to clusters of their own type, which gives
             each element its type key (P1/124, V3/120, S7, ...); unlabelled members get an
             inferred key (~key) from the types of their size.
6. rules     unlabelled walls: along the building edge or on a strip footing -> retaining wall,
             inside the building -> core wall; column-shaped pieces drawn like walls -> wall pieces.

Outputs out/geometry_classes_vN.xlsx, .pdf (one layer per class), .json (classes and pixel boxes)
and .png, N being the next free version number, so earlier runs stay for comparison.
Every tolerance is in G (this file) or T (element_stats.py), in metres or as a share; the plan's
label wording is in plan_config.toml. No colour and no CAD layer is read to detect or classify.
Requires out/detections.json from detect_structures.py for the validation columns only.

Usage
    python geometry_classes.py "GHF_EXE_PLN_STR (1)-FONDATIONS (1).pdf" --out out
"""
import argparse
import collections
import json
import math
import re
import statistics
import sys
import tomllib
from pathlib import Path
from types import SimpleNamespace

import numpy as np
import pymupdf
from scipy.optimize import linear_sum_assignment
from shapely.geometry import LineString, MultiPoint, Point, Polygon, box
from shapely.ops import unary_union
from shapely.strtree import STRtree
from sklearn.cluster import AgglomerativeClustering
from sklearn.metrics import silhouette_score

import detect_structures as ds
import element_stats as es

G = SimpleNamespace(
    area_buffer=15.0,          # objects farther than this from any solid fill are outside the drawing
    band_max=1.5,              # widest band considered (m)
    band_bin=0.01,
    band_peak_min_m=15.0,      # a band width needs this much paired line length to count
    band_peak_share=0.01,
    band_tol=0.015,
    band_empty=0.3,            # empty band: at most this share filled
    band_holds=(0.1, 0.9),     # band holding a solid: this share filled (the top only excludes a band that is the wall itself)
    hatch_spread=0.5,          # a hatched band has its strokes along most of its length (a column may fill one end)
    band_holding=0.15,         # a band family holding solids on average at least this share is a strip, not a beam
    outline_max_area=60.0,     # m2; larger closed outlines are frames, not elements
    stamp_diameter=(0.8, 3.0),
    label_radius=1.5,
    unmatched_cost=3.5,        # a label left unmatched costs as much as one 1.75 m away
    align_weight=1.0,
    long_object=6.0,           # objects longer than this can carry one label per this many metres
    name_min_votes=5,           # a name needs 5 votes, the top one at least 1.5 x the runner-up,
    name_margin=1.5,
    name_small=(2, 0.3),        # ...or a top vote of at least 2 that beats the runner-up, with 30% of the members labelled
    name_min_labelled=0.1,      # and labels on at least this share of the cluster's members
    k_range=(6, 14),            # deliberately more clusters than classes: naming merges them back
    joint_min_len=15.0,         # an expansion joint: two long lines (m) ...
    joint_gap=(0.03, 0.16),     # ... this close together (m), crossing the building footprint
    region_dominant=0.7,        # a region is one building if this share of its labels agree
    joint_margin=1.0,           # objects closer than this (m) to a region's edge (a joint) are not told its building
    affinity_min=1.0,           # a label type votes only for its home drawing kind (the one it lands on most)
    similar_max=1.6,            # an unnamed cluster takes the name of a named one closer than this (feature space)
    unnamed_penalty=1.0,        # second pass: a label may still go to an unnamed cluster, at this extra cost
    sure_distance=0.3,          # labels this close to their object fix the size of their type
    size_penalty=1.5,           # second pass: object not of the size its label's type has elsewhere
    region_penalty=1.5,         # second pass: object in a joint region of another building
    size_tol=0.03,              # two sizes (m) are the same within this: drafting and fill-outline precision
    width_tol=0.02,             # two thicknesses (m) are the same within this (thin sides are measured more tightly)
    touch_m=0.02,               # two objects closer than this (m) touch
    parallel_deg=15.0,          # two directions are parallel within this many degrees
    distance_unit=0.5,          # matching cost: one unit per this many metres between label and object
    footprint_grow=1.0,         # the building footprint: solids grown by this (m) ...
    footprint_close=3.0,        # ... and gaps up to twice this (m) closed
    edge_margin=0.6,            # a wall within footprint_grow + this (m) of the grown footprint's edge, parallel to it, is on the edge
    pen_share=0.1,              # a pen used by at least this share of the labelled columns is a column pen
    # drawing precision (m): how close counts as touching, lying on, inside or next to
    on_border=0.015,            # a line this close to a shape's border runs along that border
    inside_margin=0.005,        # a line lies between two others only if at least this far inside them
    stroke_reach=0.01,          # a hatch stroke may reach this far past the band it hatches
    fuse=0.01,                  # pieces this close are fused into one shape
    grow=0.03,                  # shapes are grown by this before testing whether they cover something
    near=0.05,                  # a band, hatch or label this close to an object is next to it
    on_region=0.1,              # an object whose centre is this close to a stamped region is on it
    dot_max=0.1,                # marks smaller than this are dots of a pattern
    # the smallest things considered (m, m2)
    line_min=0.5,               # band edges: lines at least this long ...
    overlap_min=0.5,            # ... running side by side over at least this length
    band_min=0.1,               # the narrowest band
    stroke_max=1.0,             # hatch strokes inside a band are shorter than this
    hatch_min_len=0.3,          # hatched areas shorter than this are ignored
    fill_min_area=0.02,         # m2; smaller loose fills are ignored
    bent_fill_min_area=0.5,     # m2; smaller bent fills are not split into straight stretches
    region_min_area=20.0,       # m2; smaller building regions are slivers left by the joint cuts
    wall_default=0.25,          # wall thickness assumed only if no labelled wall exists to learn it from
    # raft stamps: rays from the raft's corners to the stamp, and the outline they point to
    ray_min=1.5,                # a ray is a line at least this long ...
    ray_search=0.5,             # ... found within this of the stamp circle ...
    ray_start=0.4,              # ... that starts within this of the circle ...
    ray_reach=2.0,              # ... and reaches at least this far from the stamp's centre
    trace_zone=2.0,             # the outline is traced within this of the rays' ends
    trace_snap=0.05,            # line ends closer than this meet
    trace_min_line=0.1,         # shorter lines are ignored while tracing
    trace_wall_reach=0.1,       # a wall this close to an open end of the outline closes it
    # expansion joints
    joint_reach=0.5,            # the cut follows the joint line within this ...
    joint_cut=0.1,              # ... and is this wide
)

LABEL_CLASS = {"column": "column", "wall": "shear_wall", "retaining_wall": "retaining_wall",
               "retaining_wall+strip_footing": "strip_footing",
               "footing:S": "isolated_footing", "footing:SC": "combined_footing", "grade_beam": "grade_beam",
               "strip_footing": "strip_footing", "raft": "raft", "pit": "sump_pit"}
CLASS_COLOURS = {"column": (0.0, 0.65, 0.0), "shear_wall": (0.85, 0.0, 0.0), "retaining_wall": (0.6, 0.0, 0.3),
                 "core_wall": (0.9, 0.4, 0.6),
                 "isolated_footing": (1.0, 0.45, 0.0), "combined_footing": (0.55, 0.25, 0.0),
                 "strip_footing": (0.9, 0.7, 0.0), "grade_beam": (0.1, 0.3, 0.9), "raft": (0.5, 0.0, 0.8),
                 "sump_pit": (0.3, 0.3, 0.3)}
UNNAMED_COLOUR = (0.45, 0.45, 0.45)
# families of classes that share a geometry (a wall label can name any kind of wall)
FAMILY = {"column": "column", "shear_wall": "wall", "retaining_wall": "wall", "core_wall": "wall",
          "isolated_footing": "footing", "combined_footing": "footing", "strip_footing": "strip",
          "grade_beam": "beam", "raft": "raft", "sump_pit": "pit"}


# ---------------------------------------------------------------------------
# 1. Objects
# ---------------------------------------------------------------------------
def obj(kind, g, k, **extra):
    o = ds.oriented(g, k) if g.geom_type == "Polygon" else None
    d = {"kind": kind, "g": g, "geom": g, "o": o, "feat": {}, "label": None, "cluster": None, "cls": None}
    d.update(extra)
    return d


def band_pairs(segs, k, fill_tree, fill_geoms):
    """Every pair of eligible parallel lines up to G.band_max apart, with its spacing, overlap and
    how much of the space between them is filled. Eligible: drawn stroke, solid (not dashed), not
    in the axis pen, at least 0.5 m long. The two edges of one element are drawn with the same pen,
    so a pair must share its pen width, and only a line of that pen lying between them splits it
    (lines of other elements may cross a band). A side of a closed shape pairs only with another side
    of that same shape (an element drawn as a closed long rectangle): paired with anything else it
    bounds its own shape (a footing, a pit's double box), not a band."""
    lines = [s for s in segs if s["len"] >= G.line_min and s["stroke"] and not s.get("dashed") and not s.get("axis_pen")]
    tree = STRtree([s["g"] for s in lines])
    # short strokes, to spot hatching inside a band
    short = [s for s in segs if s["len"] < G.stroke_max and s["stroke"] and not s.get("dashed") and not s.get("axis_pen")]
    short_tree = STRtree([s["g"] for s in short])
    # a line that mostly runs along a solid's border is that solid's edge, not a band edge
    for s in lines:
        near = [fill_geoms[i] for i in fill_tree.query(s["g"].buffer(G.touch_m * k))]
        border = unary_union([g.exterior for g in near if g.geom_type == "Polygon"]) if near else None
        s["on_solid_edge"] = bool(border is not None and not border.is_empty and
                                  s["g"].intersection(border.buffer(G.on_border * k)).length > 0.5 * s["g"].length)
    out = []
    for i, s in enumerate(lines):
        ang = s["ang"]
        ux, uy = es.axis(ang)
        so = es.offset(s["a"], ang)
        s_lo, s_hi = sorted((es.along(s["a"], ang), es.along(s["b"], ang)))
        between = []
        for j in tree.query(s["g"].buffer(G.band_max * k)):
            t = lines[j]
            if j == i or abs(math.sin(t["ang"] - ang)) > 0.005:
                continue
            t_mid = ((t["a"][0] + t["b"][0]) / 2, (t["a"][1] + t["b"][1]) / 2)
            to = es.offset(t_mid, ang)
            t_lo, t_hi = sorted((es.along(t["a"], ang), es.along(t["b"], ang)))
            lo, hi = max(s_lo, t_lo), min(s_hi, t_hi)
            if (hi - lo) / k < G.overlap_min:
                continue
            between.append((to, lo, hi, j))
        for to, lo, hi, j in between:
            if j <= i or lines[j]["w"] != s["w"]:
                continue
            if (s.get("closed") or lines[j].get("closed")) and s["idx"] != lines[j]["idx"]:
                continue
            sp = abs(to - so) / k
            if sp < G.band_min or sp > G.band_max:
                continue
            inside = [j2 for o2, l2, h2, j2 in between if j2 != j and min(so, to) + G.inside_margin * k < o2 < max(so, to) - G.inside_margin * k
                      and lines[j2]["w"] == s["w"] and min(h2, hi) - max(l2, lo) > 0.3 * (hi - lo)]
            inner = bool(inside)
            # a line between the edges that is not the border of a filled element belongs to another
            # element (a beam's edge): such a pair is two elements side by side, not one band holding a wall
            foreign = any(not lines[j2]["on_solid_edge"] for j2 in inside)
            mid = (so + to) / 2
            half = sp * k / 2
            P = lambda u, v: (u * ux - v * uy, u * uy + v * ux)
            poly = Polygon([P(lo, mid - half), P(hi, mid - half), P(hi, mid + half), P(lo, mid + half)])
            filled = sum(poly.intersection(fill_geoms[q]).area for q in fill_tree.query(poly)) / poly.area
            # an empty band has nothing drawn inside; one crossed by many short oblique strokes is a
            # hatched element (a hatched wall between its two border lines)
            # hatching = one family of short parallel strokes lying within the band (lines that merely
            # cross it, such as axis dashes or dimension ticks, run on beyond its edges)
            inner_poly = poly.buffer(G.stroke_reach * k)
            strokes = [short[q] for q in short_tree.query(poly)
                       if inner_poly.contains(short[q]["g"]) and abs(math.sin(short[q]["ang"] - ang)) > 0.3]
            fam = collections.Counter(round(math.degrees(t["ang"]) / 5) for t in strokes)
            hatched = False
            if fam:
                top_dir, n_top = fam.most_common(1)[0]
                # ... spread along the band's length: a dense pattern in one spot (a pad sitting on a
                # beam) does not make the whole band a hatched wall
                pos = [es.along(((t["a"][0] + t["b"][0]) / 2, (t["a"][1] + t["b"][1]) / 2), ang)
                       for t in strokes if round(math.degrees(t["ang"]) / 5) == top_dir]
                hatched = n_top >= 3 * (hi - lo) / k and max(pos) - min(pos) >= G.hatch_spread * (hi - lo)
            if hatched:
                filled = max(filled, 1.0)
            out.append({"sp": sp, "lo": lo, "hi": hi, "ang": ang, "off": mid, "poly": poly, "filled": filled,
                        "adjacent": not inner, "foreign_inside": foreign,
                        "edges_on_solid": s["on_solid_edge"] or lines[j]["on_solid_edge"],
                        "lines": (s["idx"], lines[j]["idx"]), "overlap": (hi - lo) / k, "pen": s["w"],
                        "hatched": hatched})
    return out


def band_families(pairs):
    """Band widths used on this drawing: peaks of the paired-length histogram. A pair counts if its
    two lines face each other with nothing drawn between them, or hold a solid between them; pairs
    with an edge on a solid's face (the gap beside a wall, or the wall's own border) do not."""
    ps = [p for p in pairs if not p["edges_on_solid"] and p["filled"] <= G.band_holds[1]
          and (p["adjacent"] or (p["filled"] >= G.band_holds[0] and not p["foreign_inside"]))]
    fams = []
    if not ps:
        return fams
    nb = int(G.band_max / G.band_bin) + 2
    h = np.zeros(nb)
    for p in ps:
        h[int(round(p["sp"] / G.band_bin))] += p["overlap"]
    hs = np.convolve(h, [1, 1, 1], mode="same")
    total = h.sum()
    for b in range(1, nb - 1):
        if hs[b] >= hs[b - 1] and hs[b] > hs[b + 1] and hs[b] >= max(G.band_peak_min_m, G.band_peak_share * total):
            # the family's width: the median spacing of the pairs under the peak (weighted by their
            # length), not the bin's centre, which depends on where the bins happen to fall
            win = sorted((p["sp"], p["overlap"]) for p in ps if abs(p["sp"] - b * G.band_bin) <= 1.5 * G.band_bin)
            half, acc = sum(o for _, o in win) / 2, 0.0
            for sp, o in win:
                acc += o
                if acc >= half:
                    break
            w = round(sp, 3)
            if any(abs(f["width"] - w) <= G.band_tol for f in fams):
                continue
            pieces = [dict(p, width=w) for p in ps if abs(p["sp"] - w) <= G.band_tol]
            if not pieces:      # a smoothed peak between two populated bins, no pair at that width
                continue
            share = sum(p["filled"] * p["overlap"] for p in pieces) / sum(p["overlap"] for p in pieces)
            fams.append({"width": w, "paired_m": round(float(hs[b]), 1), "filled_share": round(share, 2),
                         "pieces": pieces})
    return fams


def circles(paths, k):
    """Closed curve-only paths that are round: stamps, bubbles, symbols."""
    out = []
    for p in paths:
        if p["type"] != "s" or not p["items"] or any(kind != "c" for kind, _ in p["items"]):
            continue
        x0, y0, x1, y1 = p["rect"]
        w, h = (x1 - x0) / k, (y1 - y0) / k
        if w > 0 and 0.85 < h / w < 1.15:
            out.append((Point((x0 + x1) / 2, (y0 + y1) / 2), w * k / 2))
    return out


def trace_outline(corners, stamp, segs, k, rays, solids=()):
    """The closed outline around a stamp: lines joined end to end, closed into a polygon. Lines of
    other elements cross an outline rather than meet it end to end, so they stay out. First the
    lines at the ray ends; if those do not close around the stamp (rays are drafting aids and can
    point to an older or unclosed box), any other closed chain around the stamp of about the rays'
    extent. Returns None if there is none."""
    from shapely import set_precision
    from shapely.ops import polygonize
    hull = MultiPoint(corners).convex_hull
    zone = hull.buffer(G.trace_zone * k)
    snap = G.trace_snap * k
    cand = [s for s in segs if s["stroke"] and not s.get("dashed") and not s.get("axis_pen") and s["len"] >= G.trace_min_line
            and s["g"].intersects(zone) and not any(s is r for r in rays)]
    key = lambda pt: (round(pt[0] / snap), round(pt[1] / snap))
    nodes = collections.defaultdict(list)
    for i, s in enumerate(cand):
        nodes[key(s["a"])].append(i)
        nodes[key(s["b"])].append(i)

    def component(start):
        seen, stack = set(start), list(start)
        while stack:
            i = stack.pop()
            for end in (cand[i]["a"], cand[i]["b"]):
                kx, ky = key(end)
                for dx in (-1, 0, 1):
                    for dy in (-1, 0, 1):
                        for j in nodes.get((kx + dx, ky + dy), []):
                            if j not in seen:
                                seen.add(j)
                                stack.append(j)
        return seen

    def loose_end_closures(ids):
        """An outline with loose ends (a side not drawn, e.g. where it meets a wall) is closed by
        joining its loose ends in nearest pairs, if each closing side is short against the outline."""
        deg = collections.Counter()
        pos = {}
        for i in ids:
            for pt in (cand[i]["a"], cand[i]["b"]):
                deg[key(pt)] += 1
                pos[key(pt)] = pt
        loose = [pos[n] for n, d in deg.items() if d == 1]
        extra, used = [], set()
        limit = 0.5 * math.sqrt(hull.area)
        pairs = sorted((math.dist(a, b), i, j) for i, a in enumerate(loose) for j, b in enumerate(loose) if i < j)
        for dd, i, j in pairs:
            if i in used or j in used or dd > limit:
                continue
            used |= {i, j}
            extra.append(LineString([loose[i], loose[j]]))
        return extra

    def closed_region(ids, extra=()):
        net = set_precision(unary_union([cand[i]["g"] for i in ids] + list(extra)), snap / 2)
        faces = list(polygonize(unary_union(net)))
        if not faces:
            return None
        u = unary_union(faces)
        g = next((g for g in getattr(u, "geoms", [u]) if g.contains(stamp)), None)
        if g is None or not 0.6 * hull.area <= g.area <= 1.6 * hull.area:
            return None
        return Polygon(g.exterior.coords)

    def closed_by_walls(ids):
        """An outline left open where it meets a wall is closed by that wall: the walls its open
        ends run into become part of the boundary; the region is the enclosed area around the stamp."""
        lines = [cand[i]["g"] for i in ids]
        ends = [Point(pt) for i in ids for pt in (cand[i]["a"], cand[i]["b"])]
        walls = [g for g in solids if g.intersects(zone) and any(g.distance(e) < G.trace_wall_reach * k for e in ends)]
        if not walls:
            return None
        barrier = unary_union([l.buffer(G.grow * k) for l in lines] + [w.buffer(G.touch_m * k) for w in walls])
        free = zone.difference(barrier)
        room = next((g for g in getattr(free, "geoms", [free]) if g.contains(stamp)), None)
        if room is None or room.boundary.intersects(zone.boundary):
            return None                         # not closed: leaks out of the search zone
        g = Polygon(room.exterior.coords).buffer(G.grow * k)
        return g if 0.6 * hull.area <= g.area <= 1.6 * hull.area else None

    start = {i for c in corners for dx in (-1, 0, 1) for dy in (-1, 0, 1)
             for i in nodes.get((key(c)[0] + dx, key(c)[1] + dy), [])}
    if start:
        comp = component(start)
        g = closed_region(comp) or closed_region(comp, loose_end_closures(comp)) or closed_by_walls(comp)
        if g is not None:
            return g
    done, best = set(), None
    for i in range(len(cand)):
        if i in done:
            continue
        comp = component({i})
        done |= comp
        g = closed_region(comp) if len(comp) > 1 else None
        if g is not None and (best is None or abs(g.area - hull.area) < abs(best.area - hull.area)):
            best = g
    return best


def stamp_rafts(paths, segs, k, solids=()):
    """Circle stamps with at least three straight rays from far away ending on the circle: the
    rays come from the corners of the element the stamp names (rafts here). Its extent is the
    outline those rays end on (not always a rectangle); the rays' convex hull if that fails."""
    long_ = [s for s in segs if s["len"] >= G.ray_min and not s.get("dashed")]
    tree = STRtree([s["g"] for s in long_])
    out, seen = [], []
    for c, r in circles(paths, k):
        if any(c.distance(q) < r for q in seen):
            continue
        corners, dirs, rays = [], set(), []
        for i in tree.query(c.buffer(r + G.ray_search * k)):
            s = long_[i]
            ux, uy = es.axis(s["ang"])
            perp = abs(-(c.x - s["a"][0]) * uy + (c.y - s["a"][1]) * ux)
            da, db = math.dist(s["a"], (c.x, c.y)), math.dist(s["b"], (c.x, c.y))
            if perp < 0.3 * r and min(da, db) < r + G.ray_start * k and max(da, db) > G.ray_reach * k:
                far = s["b"] if db > da else s["a"]
                corners.append(far)
                rays.append(s)
                dirs.add(round(math.degrees(math.atan2(far[1] - c.y, far[0] - c.x)) / 20))
        if len(dirs) >= 3:
            region = trace_outline(corners, c, segs, k, rays, solids)
            traced = region is not None
            if region is None:
                region = MultiPoint(corners).convex_hull
            if region.geom_type == "Polygon":
                seen.append(c)
                out.append(obj("stamped_region", region, k, stamp=c, rays=len(corners), traced=traced))
    return out


def objects(paths, texts, k):
    tb = es.text_index(texts)
    segs = es.straight_segments(paths, k, tb)
    es.mark_dashed(segs, k)
    log = {}
    rects, dups, others = es.filled_objects(paths, k)
    solids = [r["g"] for r in rects]
    # walls drawn around corners as one object: fills no thicker than twice the usual solid width
    thin = 2 * statistics.median(r["o"]["width_m"] for r in rects) if rects else 0.4
    for g in others:
        area = g.area / k ** 2
        if area >= G.bent_fill_min_area and g.geom_type == "Polygon" and 2 * area / (g.length / k) <= thin:
            for piece in es.straight_stretches(g, k):
                over = [s for s in solids if piece.intersection(s).area > 0]
                if not any(piece.intersection(s).area > 0.5 * piece.area for s in over):
                    solids.append(piece)
                    continue
                # the same wall drawn twice, but this drawing runs on past the other one: keep the
                # stretch that sticks out (as long as it is a stretch of the wall, not a sliver)
                w = ds.oriented(piece, k)["width_m"]
                rest = piece.difference(unary_union(over))
                for part in getattr(rest, "geoms", [rest]):
                    if part.geom_type == "Polygon" and not part.is_empty:
                        po = ds.oriented(part, k)
                        if po["width_m"] >= 0.5 * w and po["length_m"] >= w and po["rectness"] > es.T.rectness:
                            solids.append(part)
    log["solids"] = len(solids)
    out = [obj("solid", g, k) for g in solids]
    # the pen of the outline drawn around a fill, if any (how the object is drawn): the stroke width
    # of the lines running along most of its border, whether drawn as one closed ring or as pieces
    # Only the object's own lines count: lines lying (almost) entirely along its border, so the
    # outline of a neighbouring wall that runs on past it does not.
    strokes = [(ls, round(p["width"], 2)) for p in paths if p["type"] == "s" for ls in ds.path_lines(p)]
    stree = STRtree([ls for ls, _ in strokes])
    for o in out:
        border = o["g"].exterior
        band = border.buffer(G.on_border * k)
        along = collections.Counter()
        for i in stree.query(band):
            ls, w = strokes[i]
            inside = ls.intersection(band).length
            if ls.length > 0 and inside >= 0.9 * ls.length:
                along[w] += inside
        pen, cover = max(along.items(), key=lambda kv: kv[1]) if along else (0.0, 0.0)
        o["outline_pen"] = pen if cover >= 0.75 * border.length else 0.0
    top_pen = max((o["outline_pen"] for o in out), default=0) or 1.0
    for o in out:
        o["outline_pen_rel"] = o["outline_pen"] / top_pen
    for h in es.hatched_regions(segs, k):
        if h["length"] and h["length"] >= G.hatch_min_len:
            out.append(obj("cross_hatch" if h["cross"] else "single_hatch", h["g"], k, hatch=h))
    rafts = stamp_rafts(paths, segs, k, solids)
    out += rafts
    # what fills a band: solid fills and hatched areas
    fill_geoms = [o["g"] for o in out if o["kind"] in ("solid", "cross_hatch")] + [g for g in others if g.area > G.fill_min_area * k * k]
    fill_tree = STRtree(fill_geoms)
    solid_tree = STRtree([o["g"] for o in out if o["kind"] == "solid"])
    solid_list = [o for o in out if o["kind"] == "solid"]
    walls_u = unary_union([o["g"] for o in out if o["kind"] in ("solid", "cross_hatch")] +
                          [g for g in others if g.area > G.fill_min_area * k * k]).buffer(G.grow * k)
    for c in es.closed_outlines(paths, k, tb):
        g = c["g"]
        area = g.area / k ** 2
        if area > G.outline_max_area:
            continue
        if any(s["g"].intersection(g).area / s["g"].union(g).area > 0.8 for s in (solid_list[i] for i in solid_tree.query(g))):
            continue                                   # the border of a solid, not an object of its own
        if any(r["g"].intersection(g).area > 0.8 * g.area and g.area > 0.8 * r["g"].area for r in rafts):
            continue
        if walls_u.intersection(g).area > 0.8 * g.area:
            continue                                   # the border of a wall (fill or hatch)
        out.append(obj("outline", g, k))
    pairs = band_pairs(segs, k, fill_tree, fill_geoms)
    # a hatched wall's straight legs: its two border lines with the hatch between them. A bent
    # hatched region (U, L, T of walls) is replaced by its legs, each one straight.
    legs = [dict(p, width=round(p["sp"], 2)) for p in pairs if p["hatched"] and p["adjacent"]]
    runs = es.band_runs(legs, k) if legs else []
    hatch_objs = [o for o in out if o["kind"] == "cross_hatch"]
    for h in hatch_objs:
        mine = [r for r in runs if r["poly"].intersection(h["g"]).area > 0.5 * r["poly"].area]
        if not mine or h["o"]["rectness"] > 0.9:
            continue
        mine = sorted(mine, key=lambda r: -r["length"])
        kept_legs = []
        for r in mine:
            if all(r["poly"].intersection(q).area < 0.5 * r["poly"].area for q in kept_legs):
                kept_legs.append(r["poly"])
        out.remove(h)
        for leg in kept_legs:
            hh = dict(h["hatch"], thickness=min(ds.oriented(leg, k)["width_m"], 1.0), length=ds.oriented(leg, k)["length_m"])
            out.append(obj("cross_hatch", leg, k, hatch=hh, leg_of_region=True))
    log["hatched_legs"] = sum(1 for o in out if o.get("leg_of_region"))
    # one wall drawn twice, filled and hatched over the same stretch: the hatched object is that
    # solid (same direction and thickness, overlapping over most of the smaller one); the solid
    # takes the length of both, since the two drawings may stop at different points
    merged = 0
    for h in [o for o in out if o["kind"] in ("cross_hatch", "single_hatch") and o["o"]]:
        for s in (solid_list[i] for i in solid_tree.query(h["g"])):
            if not s["o"]:
                continue
            skew = abs((h["o"]["angle_deg"] - s["o"]["angle_deg"] + 90) % 180 - 90)
            same_w = abs(h["o"]["width_m"] - s["o"]["width_m"]) <= G.size_tol
            if skew <= G.parallel_deg and same_w and \
                    h["g"].intersection(s["g"]).area > 0.5 * min(h["g"].area, s["g"].area):
                u = h["o"]["obb"] and unary_union([s["g"], Polygon(h["o"]["obb"])])
                u = u if u.geom_type == "Polygon" else u.buffer(G.fuse * k).buffer(-G.fuse * k)
                if u.geom_type == "Polygon":
                    s.update(g=u, geom=u, o=ds.oriented(u, k), drawn_twice=True)
                    out.remove(h)
                    merged += 1
                    break
    log["hatched_and_filled"] = merged
    # an outline drawn around a wall that was found by its fill or hatch is that wall's border, not an
    # object of its own (checked again here, now that bent hatched regions are split into legs)
    walls_now = [o for o in out if o["kind"] in ("solid", "cross_hatch", "single_hatch")]
    wtree = STRtree([o["g"] for o in walls_now])
    borders = []
    for o in out:
        if o["kind"] != "outline":
            continue
        near = [walls_now[i] for i in wtree.query(o["g"])]
        if near and unary_union([w["g"] for w in near]).buffer(G.grow * k).intersection(o["g"]).area > 0.8 * o["g"].area:
            borders.append(o)
            continue
        # a hatched wall lying inside the outline, parallel to it (its hatch may stop short of the
        # border, e.g. where a column fills one end): the wall takes the outline's exact shape
        if not o["o"]:
            continue
        grown = o["g"].buffer(G.grow * k)
        inside = [w for w in near if w["kind"] in ("cross_hatch", "single_hatch") and w["o"]
                  and w["g"].intersection(grown).area >= 0.8 * w["g"].area
                  and abs((w["o"]["angle_deg"] - o["o"]["angle_deg"] + 90) % 180 - 90) <= G.parallel_deg]
        if inside:
            w = max(inside, key=lambda w: w["g"].area)
            w.update(g=o["g"], geom=o["g"], o=o["o"], border_shape=True)
            borders.append(o)
    out = [o for o in out if not any(o is b for b in borders)]
    log["wall_borders_dropped"] = len(borders)
    fams = band_families(pairs)
    log["band_families"] = [{"width": f["width"], "paired_m": f["paired_m"], "filled_share": f["filled_share"]} for f in fams]
    outl = [o["g"] for o in out if o["kind"] == "outline"]
    outl_tree = STRtree(outl) if outl else None
    for f in fams:
        for run in es.band_runs(f["pieces"], k):
            # the element as drawn: the whole run. Its net length leaves out the outlines it crosses
            # (a beam through a footing is footing concrete there).
            crossed = unary_union([outl[i] for i in outl_tree.query(run["poly"])]) if outl_tree is not None else None
            net = run["axis"].difference(crossed).length / k if crossed is not None and not crossed.is_empty else run["length"]
            out.append(obj("band", run["poly"], k, band=f["width"], span=run["length"], net=net, axis=run["axis"],
                           holds=f["filled_share"] >= G.band_holding))
    # the drawing area: near the solid fills (title blocks and legends have none)
    body = unary_union([o["g"] for o in out if o["kind"] == "solid"]).buffer(G.area_buffer * k)
    inside = [o for o in out if body.contains(o["g"].centroid)]
    log["outside_drawing"] = dict(collections.Counter(o["kind"] for o in out if o not in inside))
    tiny = [ds.centre(p["rect"]) for p in paths if max(p["rect"][2] - p["rect"][0], p["rect"][3] - p["rect"][1]) < G.dot_max * k]
    return inside, segs, tiny, log


# ---------------------------------------------------------------------------
# 2. Geometric features and clusters
# ---------------------------------------------------------------------------
def perimeter_flags(objs, k):
    """Walls running along the building's outer edge: close to the footprint's boundary (walls and
    columns, grown and closed) and parallel to it there."""
    from shapely.ops import nearest_points
    sol = [o["g"] for o in objs if o["kind"] in ("solid", "cross_hatch")]
    body = unary_union([g.buffer(G.footprint_grow * k) for g in sol]).buffer(G.footprint_close * k).buffer(-G.footprint_close * k)
    rings = [r for poly in getattr(body, "geoms", [body]) for r in [poly.exterior] + list(poly.interiors)]
    edges = [LineString([a, b]) for r in rings for a, b in zip(r.coords, list(r.coords)[1:])]
    etree = STRtree(edges)
    for o in objs:
        o["feat"]["on_perimeter"] = 0
        if o["kind"] not in ("solid", "cross_hatch", "outline") or not o["o"]:
            continue
        c = o["g"].centroid
        reach = (G.footprint_grow + G.edge_margin) * k        # the grown edge lies footprint_grow outside the walls
        near = [edges[i] for i in etree.query(c.buffer(reach)) if edges[i].distance(c) <= reach]
        if not near:
            continue
        e = min(near, key=lambda e: e.distance(c))
        (ax, ay), (bx, by) = e.coords[0], e.coords[-1]
        ea = math.atan2(by - ay, bx - ax)
        if abs(math.sin(ea - math.radians(o["o"]["angle_deg"]))) < math.sin(math.radians(G.parallel_deg)):
            o["feat"]["on_perimeter"] = 1


def features(objs, tiny, k, circ=()):
    solids = [o for o in objs if o["kind"] == "solid"]
    outlines = [o for o in objs if o["kind"] == "outline"]
    hold = [o for o in objs if o["kind"] == "band" and o["holds"]]
    empty = [o for o in objs if o["kind"] == "band" and not o["holds"]]
    stamped = [o for o in objs if o["kind"] == "stamped_region"]
    hatch = [o for o in objs if o["kind"] == "cross_hatch"]
    hold_u = unary_union([o["g"] for o in hold]) if hold else None
    tiny_tree = STRtree(tiny)
    s_tree, o_tree, e_tree = STRtree([o["g"] for o in solids]), STRtree([o["g"] for o in outlines]), STRtree([o["g"] for o in empty])
    for o in objs:
        g, f = o["g"], o["feat"]
        oo = o["o"]
        L, W = (oo["length_m"], oo["width_m"]) if oo else (0.0, 0.0)
        if o["kind"] == "band":
            L, W = o["span"], o["band"]
        if o["kind"] == "cross_hatch" and o["hatch"]["thickness"]:
            L, W = o["hatch"]["length"], o["hatch"]["thickness"]
        f.update(length=L, width=W, aspect=L / W if W else 0.0, area=g.area / k ** 2)
        f["solids_inside"] = sum(1 for i in s_tree.query(g) if solids[i] is not o
                                 and g.buffer(G.touch_m * k).intersection(solids[i]["g"]).area > 0.5 * solids[i]["g"].area)
        f["dots_inside"] = sum(1 for i in tiny_tree.query(g) if g.contains(tiny[i]))
        f["touches_solid"] = sum(1 for i in s_tree.query(g.buffer(G.touch_m * k)) if solids[i] is not o
                                 and solids[i]["g"].distance(g) < G.touch_m * k)
        f["inside_outline"] = int(any(outlines[i] is not o and outlines[i]["g"].buffer(G.touch_m * k).intersection(g).area > 0.5 * g.area
                                      for i in o_tree.query(g)))
        f["in_band_holding_solid"] = int(hold_u is not None and o["kind"] != "band"
                                         and hold_u.intersection(g).area > 0.5 * g.area)
        f["empty_bands_touching"] = sum(1 for i in e_tree.query(g.buffer(G.near * k)) if empty[i] is not o)
        f["on_stamped_region"] = int(any(r["g"].buffer(G.on_region * k).contains(g.centroid) for r in stamped if r is not o))
        f["touches_cross_hatch"] = int(any(h["g"].distance(g) < G.near * k for h in hatch if h is not o))
        f["outline_pen"] = o.get("outline_pen_rel", 0.0)
        f["stamp_inside"] = int(o["kind"] == "outline" and any(g.contains(c) and c.buffer(r).within(g) for c, r in circ))
    perimeter_flags(objs, k)


# drawing kind -> (feature, weight). Shape features are standardised; context flags are 0/1.
# ec2_column is the EN 1992-1-1 5.3.1(7) section test (long side <= 4 x short side): geometry,
# not a label, and the only way to split columns from walls, which leave no natural size gap.
FEATURES = {
    "solid": [("log_length", 1.0), ("log_width", 1.0), ("log_aspect", 1.0), ("ec2_column", 2.0),
              ("touches_solid", 3.0), ("inside_outline", 0.7), ("in_band_holding_solid", 1.0),
              ("on_stamped_region", 0.7), ("outline_pen", 3.0), ("on_perimeter", 2.0)],
    "outline": [("log_length", 1.0), ("log_width", 1.0), ("solids_inside", 3.0), ("log_dots", 1.0),
                ("empty_bands_touching", 0.5)],
}


def matrix(objs, spec):
    rows = []
    for o in objs:
        f = o["feat"]
        L, W = round(f["length"] * 100), round(f["width"] * 100)
        v = {"log_length": math.log(max(f["length"], 0.05)), "log_width": math.log(max(f["width"], 0.05)),
             "log_aspect": math.log(max(f["aspect"], 1.0)), "ec2_column": 1.0 if W and L <= es.T.ec2_ratio * W else 0.0,
             "touches_solid": min(f["touches_solid"], 2) / 2,
             "inside_outline": f["inside_outline"], "in_band_holding_solid": f["in_band_holding_solid"],
             "on_stamped_region": f["on_stamped_region"], "solids_inside": min(f["solids_inside"], 3) / 3,
             "log_dots": min(math.log1p(f["dots_inside"]) / 4, 1.0),
             "empty_bands_touching": min(f["empty_bands_touching"], 3) / 3, "stamp_inside": f.get("stamp_inside", 0),
             "outline_pen": f.get("outline_pen", 0.0), "on_perimeter": f.get("on_perimeter", 0)}
        rows.append([v[n] for n, _ in spec])
    X = np.array(rows, dtype=float)
    for j, (n, w) in enumerate(spec):
        if n.startswith("log_") and n != "log_dots" and X[:, j].std() > 0:
            X[:, j] = (X[:, j] - X[:, j].mean()) / X[:, j].std()
        X[:, j] *= w
    return X


def cluster(objs):
    """Clusters per drawing kind. Solids and outlines: agglomerative (Ward) clustering on weighted
    geometric features, deliberately into more clusters than there are classes (the count with the
    best silhouette score within G.k_range). Bands: one cluster per band width found in the drawing.
    Hatches and stamped regions: one cluster per kind."""
    info = []
    by_kind = collections.defaultdict(list)
    for o in objs:
        by_kind[o["kind"]].append(o)
    cid = 0
    for kind, os_ in sorted(by_kind.items()):
        if kind in FEATURES and len(os_) > G.k_range[0] * 3:
            X = matrix(os_, FEATURES[kind])
            best = None
            for n in range(G.k_range[0], min(G.k_range[1], len(os_) - 1) + 1):
                lab = AgglomerativeClustering(n_clusters=n, linkage="ward").fit_predict(X)
                sc = silhouette_score(X, lab)
                if best is None or sc > best[0]:
                    best = (sc, n, lab)
            sc, n, lab = best
            for i, o in enumerate(os_):
                o["x"] = X[i]
            for j in range(n):
                idx = [i for i, l in enumerate(lab) if l == j]
                for i in idx:
                    os_[i]["cluster"] = cid
                info.append({"id": cid, "kind": kind, "n": len(idx), "silhouette": round(float(sc), 3),
                             "centroid": X[idx].mean(axis=0)})
                cid += 1
        elif kind == "band":
            for w in sorted({o["band"] for o in os_}):
                members = [o for o in os_ if o["band"] == w]
                for o in members:
                    o["cluster"] = cid
                info.append({"id": cid, "kind": f"band {w:.2f} m", "n": len(members), "silhouette": None, "centroid": None})
                cid += 1
        else:
            for o in os_:
                o["cluster"] = cid
            info.append({"id": cid, "kind": kind, "n": len(os_), "silhouette": None, "centroid": None})
            cid += 1
    return info


# ---------------------------------------------------------------------------
# 3. Labels and matching
# ---------------------------------------------------------------------------
def read_labels(texts, paths, cfg, k, tb):
    labels = es.all_labels(texts, cfg)
    for l in labels:
        if l["kind"] == "footing":
            l["kind"] = "footing:" + l["groups"]["series"]
    # labels naming a region (raft stamp, pit note): their wording comes from the config
    rr, pr = cfg["region_labels"]["raft"], cfg["region_labels"]["pit"]
    nums = [t for t in texts if re.fullmatch(rr["number"], t["text"])]
    for t in texts:
        if re.search(rr["pattern"], t["text"]) and not (rr.get("exclude") and re.search(rr["exclude"], t["text"])):
            c = ds.centre(t["rect"])
            m = re.search(rr["thickness"], t["text"])
            thick = m.group(1) if m else None
            if thick is None and nums:     # the stamp's number sits right under the word
                n = min(nums, key=lambda n: ds.centre(n["rect"]).distance(c))
                thick = n["text"] if ds.centre(n["rect"]).distance(c) < G.label_radius * k else None
            labels.append({"kind": "raft", "raw": t["text"], "c": c, "key": f"{rr['key']} {thick}" if thick else rr["key"],
                           "groups": {}, "building": None, "dims": None})
        if re.search(pr["pattern"], t["text"]):
            labels.append({"kind": "pit", "raw": t["text"], "c": ds.centre(t["rect"]), "key": pr["key"],
                           "groups": {}, "building": None, "dims": None})
    # one text naming two stacked elements ('VP20-SF1 80x25': the wall and the strip under it)
    merged, used = [], set()
    for a in labels:
        if a["kind"] != "retaining_wall" or id(a) in used:
            continue
        b = next((b for b in labels if b["kind"] == "strip_footing" and b["raw"] == a["raw"] and b["c"].equals(a["c"])), None)
        if b is not None:
            used.update((id(a), id(b)))
            merged.append(dict(a, kind="retaining_wall+strip_footing", key=b["key"], parts={"retaining_wall": a, "strip_footing": b}))
    labels = [l for l in labels if id(l) not in used] + merged
    # text direction from the glyph strokes inside each label's box
    glyphs = [s for p in paths if p["type"] == "s" for s in ds.path_lines(p)]
    gtree = STRtree(glyphs)
    for l in labels:
        t = next((t for t in texts if t["text"] == l["raw"] and ds.centre(t["rect"]).equals(l["c"])), None)
        l["dir"] = None
        if t is None:
            continue
        b = box(*t["rect"])
        # a glyph stroke lies inside the text's box and is no longer than the text is tall
        tall = 1.2 * min(t["rect"][2] - t["rect"][0], t["rect"][3] - t["rect"][1])
        pts = [c for i in gtree.query(b) if b.contains(glyphs[i]) and glyphs[i].length <= tall for c in glyphs[i].coords]
        if len(pts) >= 6:
            mrr = MultiPoint(pts).minimum_rotated_rectangle
            if mrr.geom_type == "Polygon":
                q = list(mrr.exterior.coords)
                a, bb = math.dist(q[0], q[1]), math.dist(q[1], q[2])
                p0, p1 = (q[0], q[1]) if a >= bb else (q[1], q[2])
                if max(a, bb) > 1.8 * min(a, bb):
                    l["dir"] = math.atan2(p1[1] - p0[1], p1[0] - p0[0]) % math.pi
                    l["text_len"] = max(a, bb) / k
    return labels


def label_distance(l, o, k):
    """Metres from a label to an object: to the solid itself; to the edge of a closed outline (a
    label deep inside a footing is about what stands in it); to the stamp of a stamped region."""
    if o["kind"] == "outline":
        return 0.0 if o["g"].exterior.distance(l["c"]) < G.near * k else o["g"].exterior.distance(l["c"]) / k
    if o["kind"] == "stamped_region":
        return o["stamp"].distance(l["c"]) / k
    return 0.0 if o["g"].contains(l["c"]) else o["g"].distance(l["c"]) / k


def match(labels, objs, k, penalty=None):
    """One global assignment of labels to objects. Cost = distance (in half-metres) + misalignment
    between the text and a long object's axis (+ penalty(label, object), None = not allowed);
    long objects offer one slot per G.long_object metres; each label may instead stay unmatched
    at G.unmatched_cost."""
    tree = STRtree([o["g"].buffer(0) for o in objs])
    slots = []
    for i, o in enumerate(objs):
        L = o["feat"]["length"]
        slots += [i] * (1 + int(L // G.long_object) if L > G.long_object else 1)
    slot_of = collections.defaultdict(list)
    for s, i in enumerate(slots):
        slot_of[i].append(s)
    BIG = 1e6
    C = np.full((len(labels), len(slots) + len(labels)), BIG)
    detail = {}
    for li, l in enumerate(labels):
        C[li, len(slots) + li] = G.unmatched_cost
        for i in tree.query(l["c"].buffer(G.label_radius * k)):
            o = objs[i]
            extra = 0.0
            if penalty is not None:
                extra = penalty(l, o)
                if extra is None:
                    continue
            d = label_distance(l, o, k)
            if d > G.label_radius:
                continue
            mis = 0.0
            # a text is laid along an object only if the object is at least as long as the text
            if l["dir"] is not None and o["o"] and o["feat"]["aspect"] >= 2 and o["feat"]["length"] >= l.get("text_len", 0):
                mis = abs(math.sin(l["dir"] - math.radians(o["o"]["angle_deg"])))
            cost = d / G.distance_unit + G.align_weight * mis + extra
            for s in slot_of[i]:
                C[li, s] = cost
            detail[(li, i)] = (round(d, 2), round(mis, 2))
    rows, cols = linear_sum_assignment(C)
    out = {}
    for li, s in zip(rows, cols):
        if s < len(slots) and C[li, s] < BIG:
            i = slots[s]
            out[li] = (i, float(C[li, s]), detail[(li, i)])
    return out


def dims_agree(l, o):
    """A label that writes its section (LG-40x15, SF1 80x25) only speaks for a band whose width is
    one of those numbers."""
    nums = l.get("numbers") or (l.get("parts", {}).get("strip_footing") or {}).get("numbers")
    if not nums or o["kind"] != "band":
        return True
    return any(abs(o["band"] - n) <= G.size_tol for n in nums)


def compound_targets(l, i, objs):
    """(object, label kind) pairs a matched label speaks for. A 'VP..-SF..' label on a band
    names the band (strip) and the solid standing in it (wall); on a solid, only the wall."""
    if l["kind"] != "retaining_wall+strip_footing":
        return [(i, l["kind"])]
    o = objs[i]
    if o["kind"] != "band":
        return [(i, "retaining_wall")]
    inside = [j for j, q in enumerate(objs) if q["kind"] == "solid" and q["g"].intersects(o["g"])
              and o["g"].intersection(q["g"]).area > 0.5 * q["g"].area]
    out = [(i, "strip_footing")]
    if inside:
        out.append((min(inside, key=lambda j: objs[j]["g"].distance(l["c"])), "retaining_wall"))
    return out


HOME = {}     # label type -> the drawing kind(s) it lands on most, learned in name_clusters


def label_size_meaning(cfg):
    """What a label type's key says about the element's size in plan: both sides (P, V, S types:
    one type, one section), the width only (VP thickness, LG / SF sections) or nothing (a raft's
    thickness, a pit note). Read from the label grammar in the config."""
    dims = {kind: ("width" if r.get("size_match") == "width" else "both") for kind, r in cfg["labels"].items()}
    dims.update({kind: "width" for kind in cfg["bands"]})
    for kind in list(dims):
        if kind == "footing":
            dims.update({f"footing:{sr}": dims[kind] for sr in cfg["labels"]["footing"].get("series_subclass", {})})
    dims["retaining_wall+strip_footing"] = "width"
    return dims


def landing_scores(labels, objs, m, dims):
    """How strongly each label type belongs to each drawing kind, from the neutral pass. A type label
    names elements of one size, so a landing counts when its object has the size of another object
    the same type key landed on on that drawing kind (the type comes out consistent there). Label
    types whose key says nothing about the size, or with fewer size-consistent landings than it takes
    to name a cluster, count plain landings."""
    raw = collections.defaultdict(collections.Counter)
    groups = collections.defaultdict(list)
    for li, (i, cost, _) in m.items():
        l, o = labels[li], objs[i]
        raw[l["kind"]][o["kind"]] += 1
        groups[(l["kind"], o["kind"], l["key"])].append((o["feat"]["length"], o["feat"]["width"]))
    cons = collections.defaultdict(collections.Counter)
    for (kd, dk, key), sizes in groups.items():
        what = dims.get(kd)
        if what is None:
            continue
        same = lambda a, b: abs(a[1] - b[1]) <= G.size_tol and (what == "width" or abs(a[0] - b[0]) <= G.size_tol)
        cons[kd][dk] += sum(1 for j, a in enumerate(sizes) if any(same(a, b) for jj, b in enumerate(sizes) if jj != j))
    # consistency decides only with at least as much evidence as it takes to name a cluster
    return {kd: (cons[kd] if sum(cons[kd].values()) >= G.name_min_votes else raw[kd]) for kd in raw}


def name_clusters(info, labels, objs, m, dims=None):
    """A cluster takes the class of the label type most of its matched labels carry (at least
    G.name_min_votes labels, the top one at least G.name_margin x the runner-up). A cluster with too few labels
    takes the name of the nearest named cluster of the same drawing kind, if close enough in
    feature space; otherwise it stays unnamed."""
    # which drawing kinds each label type lands on, learned from the neutral pass itself (P labels on
    # solids, LG labels on bands, S labels on outlines...): a label type votes only for drawing kinds
    # it lands on often, so a stray P label next to a band does not make that band a column
    lands = landing_scores(labels, objs, m, dims or {})
    # relative to the kind it belongs to most (its home kind)
    affinity = {kd: {dk: n / max(c.values()) for dk, n in c.items()} if c and max(c.values()) else {}
                for kd, c in lands.items()}
    HOME.clear()
    HOME.update({kd: {dk for dk, a in aff.items() if a >= G.affinity_min} for kd, aff in affinity.items()})
    votes = collections.defaultdict(collections.Counter)
    for li, (i, cost, _) in m.items():
        if not dims_agree(labels[li], objs[i]):
            continue
        if affinity.get(labels[li]["kind"], {}).get(objs[i]["kind"], 0) < G.affinity_min:
            continue
        for oi, kind in compound_targets(labels[li], i, objs):
            votes[objs[oi]["cluster"]][kind] += 1
    for c in info:
        v = votes[c["id"]]
        total = sum(v.values())
        c.update(votes=dict(v), labelled=total, name=None, purity=None, named_by="")
        if total and total >= G.name_min_labelled * c["n"]:
            top = v.most_common(2)
            kind, n = top[0]
            second = top[1][1] if len(top) > 1 else 0
            if (n >= G.name_min_votes and n >= G.name_margin * second) or \
                    (n >= G.name_small[0] and n > second and total >= G.name_small[1] * c["n"]):
                c.update(name=LABEL_CLASS.get(kind), purity=n / total, named_by="label vote")
    for c in info:
        if c["name"] or c["centroid"] is None:
            continue
        named = [q for q in info if q["name"] and q["kind"] == c["kind"] and q["named_by"] == "label vote"]
        if named:
            q = min(named, key=lambda q: float(np.linalg.norm(q["centroid"] - c["centroid"])))
            dist = float(np.linalg.norm(q["centroid"] - c["centroid"]))
            v = votes[c["id"]]
            top = [kd for kd, n in v.items() if v and n == max(v.values())]
            # close enough, or the nearest named cluster's class is one of this cluster's own top votes (a tie)
            if dist <= G.similar_max or any(LABEL_CLASS.get(kd) == q["name"] for kd in top):
                c.update(name=q["name"], named_by=f"similar to cluster {q['id']} (distance {dist:.2f})")
    names = {c["id"]: c["name"] for c in info}
    how = {c["id"]: c["named_by"] for c in info}
    for o in objs:
        o["cls"] = names[o["cluster"]]
        o["named_by"] = how[o["cluster"]]
    return votes


# ---------------------------------------------------------------------------
# 4. Buildings, type keys, validation
# ---------------------------------------------------------------------------
def joints(segs, k):
    """Expansion joints: pairs of long parallel lines drawn with the same pen, a few centimetres
    apart. They separate buildings."""
    long_ = [s for s in segs if s["len"] >= G.joint_min_len and s["stroke"] and not s.get("dashed")
             and not s.get("axis_pen")]
    out = []
    for i, a in enumerate(long_):
        for b in long_[i + 1:]:
            if a["w"] != b["w"] or abs(math.sin(a["ang"] - b["ang"])) > 0.005:
                continue
            gap = abs(es.offset(((b["a"][0] + b["b"][0]) / 2, (b["a"][1] + b["b"][1]) / 2), a["ang"]) - es.offset(a["a"], a["ang"])) / k
            if not G.joint_gap[0] <= gap <= G.joint_gap[1]:
                continue
            la = sorted((es.along(a["a"], a["ang"]), es.along(a["b"], a["ang"])))
            lb = sorted((es.along(b["a"], a["ang"]), es.along(b["b"], a["ang"])))
            if min(la[1], lb[1]) - max(la[0], lb[0]) > 0.8 * min(la[1] - la[0], lb[1] - lb[0]):
                out.append(unary_union([a["g"], b["g"]]))
    return out


def building_regions(objs, labels, segs, k):
    """Split the building footprint along the expansion joints; each region takes the building
    number most labels inside it carry. Regions whose labels disagree (a boundary not drawn as a
    joint) are left to nearest-label voting within the region."""
    body = unary_union([o["g"].buffer(G.footprint_grow * k) for o in objs if o["kind"] in ("solid", "outline")])
    body = body.buffer(G.footprint_close * k).buffer(-G.footprint_close * k)
    cuts = []
    span = max(body.bounds[2] - body.bounds[0], body.bounds[3] - body.bounds[1])
    for j in joints(segs, k):
        ls = max(getattr(j, "geoms", [j]), key=lambda g: g.length)
        (ax, ay), (bx, by) = ls.coords[0], ls.coords[-1]
        L = math.hypot(bx - ax, by - ay)
        ux, uy = (bx - ax) / L, (by - ay) / L
        # extend until it leaves the footprint on both sides, keeping only the stretch through the
        # joint's own part of the building (a U-shaped plan must not be cut across its other arm)
        ext = LineString([(ax - ux * span, ay - uy * span), (bx + ux * span, by + uy * span)]).intersection(body)
        piece = next((g for g in getattr(ext, "geoms", [ext]) if g.distance(ls) < G.joint_reach * k and g.length > 0), None)
        if piece is not None:
            cuts.append(piece.buffer(G.joint_cut * k))
    regions = body.difference(unary_union(cuts)) if cuts else body
    regions = [g for g in getattr(regions, "geoms", [regions]) if g.area > G.region_min_area * k * k]
    info = []
    for g in regions:
        v = collections.Counter(l["building"] for l in labels if l.get("building") and g.contains(l["c"]))
        total = sum(v.values())
        b, n = v.most_common(1)[0] if v else (None, 0)
        info.append({"g": g, "building": b if total and n / total >= G.region_dominant else None,
                     "labels": [l for l in labels if l.get("building") and g.contains(l["c"])], "votes": dict(v)})
    return info, len(cuts)



# ---------------------------------------------------------------------------
def wanted(l):
    """The classes a label can name."""
    if l["kind"] == "retaining_wall+strip_footing":
        return {"strip_footing", "retaining_wall"}
    return {LABEL_CLASS[l["kind"]]}


def assign_types(objs, labels, m2, k, cfg, regions=None):
    """Type key of each object from its second-pass label; for an unlabelled member of a named
    class, the one type of its class (and building, where types vary per building) with its size."""
    for o in objs:
        o.update(key=None, label_text=None, label_cost=None, label_d=None, label_mis=None, label_src=None,
                 type_source="")
    for li, (i, cost, (d, mis)) in m2.items():
        l = labels[li]
        for oi, kind in compound_targets(l, i, objs):
            o = objs[oi]
            src = l["parts"][kind] if "parts" in l else l
            if o["key"] is None or cost < o["label_cost"]:
                o.update(key=src["key"], label_text=l["raw"], label_cost=cost, label_d=d, label_mis=mis,
                         label_src=src, type_source="printed label")
    blds = es.Buildings([l for l in labels if l.get("building")], k)
    for o in objs:
        own = (o["label_src"] or {}).get("groups", {}).get("bld") if o["label_src"] else None
        c = o["g"].centroid
        reg = next((r for r in regions or [] if r["g"].contains(c)), None)
        if reg is not None and reg["building"]:
            where, src = reg["building"], "region between joints"
        elif reg is not None and reg["labels"]:
            where, conf = es.Buildings(reg["labels"], k).of(c)
            src = "nearest labels in its region"
        else:
            where, conf = blds.of(c)
            src = "nearest labels"
        if own:
            o["building"], o["building_source"] = own, "own label" + ("" if where in (None, own) else f" (region says {where})")
        else:
            o["building"], o["building_source"] = where, src
    rule_of = {"column": cfg["labels"]["column"], "shear_wall": cfg["labels"]["wall"],
               "retaining_wall": cfg["labels"]["retaining_wall"], "isolated_footing": cfg["labels"]["footing"],
               "combined_footing": cfg["labels"]["footing"]}

    def fits_size(o, qs, rule):
        Ls = statistics.median(q["feat"]["length"] for q in qs)
        Ws = statistics.median(q["feat"]["width"] for q in qs)
        ok_w = abs(o["feat"]["width"] - Ws) <= G.width_tol
        return ok_w if rule.get("size_match") == "width" else ok_w and abs(o["feat"]["length"] - Ls) <= G.size_tol

    # a printed label is checked against the size its type has on the other elements carrying it;
    # when it does not fit, the types of the building region the element stands in are tried, then
    # the other types of the label's building; otherwise '?'. A type seen on only two elements that
    # disagree cannot say which is right: both are flagged and keep their label.
    printed = collections.defaultdict(list)
    kind_cls = {"column": "column", "wall": "shear_wall", "retaining_wall": "retaining_wall"}
    for o in objs:
        src = o.get("label_src")
        if o["type_source"] == "printed label" and src and src.get("kind") in kind_cls:
            printed[(kind_cls[src["kind"]], o["key"])].append(o)
    retyped = []
    for (cls, key), group in printed.items():
        rule = rule_of[cls]
        for o in group:
            others = [q for q in group if q is not o]
            # it fits if it has the size of any other element carrying its type
            if not others or any(fits_size(o, [q], rule) for q in others):
                continue
            # the type's size is only known if at least two of the others agree on it
            agreeing = max(([r for r in others if fits_size(r, [q], rule)] for q in others), key=len)
            if len(agreeing) < 2:
                o["type_source"] = f"printed label; size differs from the other {key} (conflict, check)"
                continue
            retyped.append((o, cls, key))
    for o, cls, key in retyped:
        rule = rule_of[cls]
        cands = []
        for bld, src in ((o.get("region_bld"), "the building region it stands in"), (o["building"], "its label's building")):
            if not bld:
                continue
            for (c2, k2), qs in printed.items():
                if c2 == cls and k2 != key and qs[0]["building"] == bld and fits_size(o, qs, rule):
                    cands.append((k2, src))
            if cands:
                break
        keys = sorted({k2 for k2, _ in cands})
        if len(keys) == 1:
            o["key"] = keys[0]
            o["type_source"] = f"printed {key} does not fit that type's size; the size fits {keys[0]} of {cands[0][1]}"
        elif keys:
            o["key"] = "|".join(keys)
            o["type_source"] = f"printed {key} does not fit that type's size; several types of {cands[0][1]} fit"
        else:
            o["key"] = None
            o["type_source"] = f"printed {key} does not fit that type's size; no type of its building fits"
    groups = collections.defaultdict(list)
    for o in objs:
        if o["key"] and o["cls"] in rule_of and o["type_source"].startswith("printed"):
            groups[(o["cls"], o["key"])].append(o)
    for o in objs:
        if o["key"] or o["cls"] not in rule_of or o["type_source"].startswith("printed"):
            continue
        rule = rule_of[o["cls"]]
        fits = []
        for (cls, key), qs in groups.items():
            if cls != o["cls"]:
                continue
            same = [q for q in qs if not rule["per_building"] or q["building"] == o["building"]]
            if not same:
                continue
            Ls = statistics.median(q["feat"]["length"] for q in same)
            Ws = statistics.median(q["feat"]["width"] for q in same)
            ok_w = abs(o["feat"]["width"] - Ws) <= G.width_tol
            ok = ok_w if rule.get("size_match") == "width" else ok_w and abs(o["feat"]["length"] - Ls) <= G.size_tol
            if ok:
                # nearest element of that type, to rank several candidates
                fits.append((min(q["g"].distance(o["g"]) for q in same) / k, key))
        where = " in its building" if rule["per_building"] else ""
        if len(fits) == 1:
            o["key"] = fits[0][1]
            o["type_source"] = f"inferred: the only {o['cls']} type{where} with this size"
        elif len(fits) > 1:
            keys = [key for _, key in sorted(fits)]
            o["key"] = "|".join(keys)
            o["type_source"] = f"inferred, ambiguous: {len(keys)} {o['cls']} types{where} have this size (nearest first)"
        else:
            o["type_source"] = f"no {o['cls']} type{where} has this size"


def validate(objs, ref_path, k):
    refs = es.load_reference(ref_path, k)
    rows = []
    for r in refs:
        if r["subclass"] in ("wall_network", "crane_foundation"):
            continue
        best, bo = 0.0, None
        for o in objs:
            if o["g"].intersects(r["g"]):
                iou = o["g"].intersection(r["g"]).area / o["g"].union(r["g"]).area
                if iou > best:
                    best, bo = iou, o
        if bo is not None and best > 0.8:
            bo["verified"] = r["subclass"]
        rows.append((r["subclass"], bo["cls"] if bo is not None and best > 0.8 else "not found"))
    return rows


# ---------------------------------------------------------------------------
# 5. Outputs
# ---------------------------------------------------------------------------
OBJ_COLS = [("ID", 6), ("Building", 9), ("Building from", 22), ("Drawing kind", 14), ("Cluster", 8), ("Class", 17),
            ("Label shown", 34), ("Class named by", 26),
            ("Type key", 14), ("Type source", 30), ("Label text", 20), ("Length (cm)", 9), ("Width (cm)", 9), ("L/W", 7),
            ("Area (m²)", 9), ("Angle (°)", 8), ("Touches solids", 8), ("Inside an outline", 8), ("In a strip band", 8),
            ("On a stamped region", 8), ("On the building edge", 8), ("Solids inside", 8), ("Dots inside", 8), ("Beams touching", 8), ("Net length (m)", 9), ("Extent", 14),
            ("bbox x0 (px)", 9), ("bbox y0 (px)", 9), ("bbox x1 (px)", 9), ("bbox y1 (px)", 9),
            ("(validation) verified class", 16), ("(validation) agrees", 10)]
VERIFIED_TO_CLASS = {"column": "column", "shear_wall": "shear_wall", "retaining_wall": "retaining_wall",
                     "isolated_footing": "isolated_footing", "combined_footing": "combined_footing",
                     "crane_footing": "isolated_footing", "mass_concrete_pad": None}


def cletter(cols, name):
    from openpyxl.utils import get_column_letter
    return get_column_letter([c for c, _ in cols].index(name) + 1)


def write_workbook(path, objs, info, labels, m1, m2, log, ver, dpi, calib):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    F = "Arial"
    hf, hfill = Font(name=F, bold=True, color="FFFFFF"), PatternFill("solid", fgColor="2F4F6F")
    body, bold, title = Font(name=F, size=10), Font(name=F, size=10, bold=True), Font(name=F, size=14, bold=True)
    note = Font(name=F, size=9, italic=True)
    wrap = Alignment(wrap_text=True, vertical="top")

    def header(ws, row, names, widths=None, col0=1):
        for i, n in enumerate(names):
            c = ws.cell(row=row, column=col0 + i, value=n)
            c.font, c.fill, c.alignment = hf, hfill, Alignment(wrap_text=True, vertical="center")
            if widths:
                ws.column_dimensions[get_column_letter(col0 + i)].width = widths[i]
        ws.row_dimensions[row].height = 30

    s = dpi / 72
    k = calib["pt_per_m"]
    wb = Workbook()
    order = {c: i for i, c in enumerate(["column", "shear_wall", "core_wall", "retaining_wall", "isolated_footing", "combined_footing",
                                         "strip_footing", "grade_beam", "raft", "sump_pit"])}
    objs = sorted(objs, key=lambda o: (str(o["building"]), order.get(o["cls"], 99), str(o["cls"]), str(o["key"])))
    for n, o in enumerate(objs, 1):
        o["id"] = n
    cinfo = {c["id"]: c for c in info}
    N = len(objs) + 1
    OE = lambda name: f"Objects!${cletter(OBJ_COLS, name)}$2:${cletter(OBJ_COLS, name)}${N}"

    # --- README ---
    ws = wb.active
    ws.title = "README"
    text = [
        ("Geometry-first classification", title), ("", body),
        ("How each element got its class", bold),
        ("1. Objects were found from geometry only: solid fills (pieces grouped by the order they are drawn in), "
         "cross-hatched regions, closed outlines, bands (pairs of parallel lines; the widths come from the drawing "
         "itself) and circle stamps with corner rays. No colour, no layer, no text.", body),
        ("2. Objects of the same drawing kind were clustered on geometric features only: length, width, the ratio of the "
         "two (log), the EN 1992-1-1 5.3.1(7) section test (long side <= 4 x short side = column-shaped), whether it "
         "touches another solid, lies inside a closed outline, stands in a strip band or on a raft, and the pen of the "
         "outline drawn around it (none / thin / thick, relative to the thickest on the sheet). For outlines: length, "
         "width, how many solids and stipple dots are inside, how many beams touch it. 'Similar' means close in these "
         "features after scaling, not the same dimensions. Deliberately more clusters than classes.", body),
        ("3. Every label was matched to one object by distance and text alignment, in one global assignment that "
         "ignores what the label says; a label may stay unmatched.", body),
        (f"4. A cluster takes the class that most of its matched labels name (>= {G.name_min_votes} votes for it, at "
         f"least {G.name_margin} x the runner-up, with labels on >= {G.name_min_labelled:.0%} of its members, or a small "
         "clear vote of 2 or more with labels on 30% of its members). A cluster without enough labels takes the name of the most "
         f"similar named cluster of the same drawing kind (feature distance <= {G.similar_max}), else stays unnamed. "
         "Unlabelled members inherit their cluster's class - that is how unlabelled walls became walls.", body),
        ("5. Labels were matched again, now only to objects of the class they name (or unnamed ones, at a cost), "
         "penalising an object whose size differs from the size its type has where a label sits right on its element, "
         "or that lies between expansion joints of another building. That gives the type keys (P1/124, V3/120, S7...).", body),
        ("6. An unlabelled member of a class takes the type of its class (and building, where sizes vary per building) "
         "that has its size: '~P1/124' when one type fits, '~V2/123|V4/123' when several share the size (nearest first). "
         "When none fits, the type source says so: for a 'column' that is a sign it is not a column.", body),
        ("7. Buildings: the footprint is cut along the expansion joints (pairs of long lines a few cm apart); a region "
         "whose labels agree is one building. Elsewhere the nearest building-numbered labels within the region decide.", body),
        ("8. A printed label whose element's size differs from the size its type has on at least two other elements "
         "is re-typed from the types of the building region the element stands in (then its label's building), else "
         "'?'. Two elements of a type that disagree are both flagged 'conflict'.", body),
        ("9. Unlabelled walls (solid or hatched): along the building edge (close and parallel to the footprint's "
         "boundary) or standing on a strip footing they are retaining walls; inside the building, when no V type has their size, core walls. A "
         "column-cluster member without a label stays a column unless it is drawn unlike the labelled columns (an "
         "outline pen none of them has) or continues a wall of the same thickness; then it is a wall piece.", body),
        ("", body),
        ("A label type only votes for its home drawing kind - the one it lands on most in the neutral pass (learned "
         "from the drawing: P and V labels land on solids, LG and SF labels on bands, S labels on outlines).", body),
        ("", body),
        ("What labels do and do not decide", bold),
        ("A single label never decides an object's class: it votes for its cluster, and its object keeps the "
         "cluster's class even if the label disagrees. Disagreeing labels are listed on 'Conflicts' - each is a "
         "wrong label, a wrongly placed label, or a geometry mistake worth a look.", body),
        ("Clusters no label votes for (cross-hatched walls, small stippled squares, several band widths) keep a "
         "descriptive geometric name; naming them is an interpretation, not something the plan states.", body),
        ("", body),
        ("Sheets", bold),
        ("Summary by building - counts and sizes per building, class and type (formulas over 'Objects').", body),
        ("Clusters - every geometric cluster, its size ranges, the label votes it received and the name it took.", body),
        ("Objects - every object with its class, type, measurements, features and pixel box (page at 150 dpi).", body),
        ("Labels - every label: what it names, where the first (neutral) pass put it, whether that agreed, and the "
         "second-pass match.", body),
        ("Conflicts - labels that disagreed with their object's cluster, and labels left unmatched.", body),
        ("Validation - geometric classes against the verified, colour-based detections.", body),
        ("", body),
        (f"Scale 1/{calib['scale_denominator']:.1f}, from {calib['inliers']} dimension texts (median error "
         f"{calib['median_error_cm']:.2f} cm). Formulas recalculate on opening; they were checked with the Python "
         "'formulas' engine.", body),
    ]
    for i, (t, f) in enumerate(text, 1):
        c = ws.cell(row=i, column=1, value=t)
        c.font, c.alignment = f, wrap
    ws.column_dimensions["A"].width = 130

    # --- Objects ---
    ws = wb.create_sheet("Objects")
    header(ws, 1, [c for c, _ in OBJ_COLS], [w for _, w in OBJ_COLS])
    Lc, Wc = cletter(OBJ_COLS, "Length (cm)"), cletter(OBJ_COLS, "Width (cm)")
    Cc, Vc = cletter(OBJ_COLS, "Class"), cletter(OBJ_COLS, "(validation) verified class")
    for r, o in enumerate(objs, 2):
        f = o["feat"]
        x0, y0, x1, y1 = o["g"].bounds
        c = cinfo[o["cluster"]]
        vals = {"ID": o["id"], "Building": o["building"] or "unknown", "Building from": o.get("building_source", ""),
                "Drawing kind": o["kind"], "Cluster": o["cluster"],
                "Class": o["cls"] or "unnamed", "Label shown": display_label(o),
                "Class named by": o.get("wall_rule") or o.get("named_by") or "no label votes",
                "Type key": o["key"] or "(none)", "Type source": o["type_source"], "Label text": o["label_text"] or "",
                "Length (cm)": round(f["length"] * 100), "Width (cm)": round(f["width"] * 100),
                "L/W": f'=IF(AND(ISNUMBER({Lc}{r}),{Wc}{r}>0),{Lc}{r}/{Wc}{r},"")',
                "Area (m²)": round(f["area"], 3), "Angle (°)": round(o["o"]["angle_deg"], 1) if o["o"] else None,
                "Touches solids": f["touches_solid"], "Inside an outline": f["inside_outline"],
                "In a strip band": f["in_band_holding_solid"], "On a stamped region": f["on_stamped_region"],
                "On the building edge": f.get("on_perimeter", 0),
                "Solids inside": f["solids_inside"], "Dots inside": f["dots_inside"], "Beams touching": f["empty_bands_touching"],
                "Net length (m)": round(o["net"], 2) if o.get("net") is not None else None,
                "Extent": ("traced outline" if o.get("traced") else "diagonals' rectangle (no closed outline drawn)")
                if o["kind"] == "stamped_region" else "",
                "bbox x0 (px)": round(x0 * s, 1), "bbox y0 (px)": round(y0 * s, 1), "bbox x1 (px)": round(x1 * s, 1),
                "bbox y1 (px)": round(y1 * s, 1), "(validation) verified class": o.get("verified", ""),
                "(validation) agrees": (f'=IF({Vc}{r}="","",IF(OR({Vc}{r}={Cc}{r},AND({Vc}{r}="crane_footing",'
                                        f'{Cc}{r}="isolated_footing")),"yes","no"))')}
        for ci, (name, _) in enumerate(OBJ_COLS, 1):
            ws.cell(row=r, column=ci, value=vals[name]).font = body
        ws.cell(row=r, column=[c for c, _ in OBJ_COLS].index("L/W") + 1).number_format = "0.00"
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(OBJ_COLS))}{N}"

    # --- Labels ---
    ws = wb.create_sheet("Labels")
    LCOLS = [("Label", 5), ("Text", 20), ("Kind", 22), ("Names class", 26), ("x (px)", 8), ("y (px)", 8),
             ("Pass 1 object", 8), ("Pass 1 cluster", 8), ("Pass 1 cluster class", 17), ("Pass 1 agrees", 9),
             ("Pass 2 object", 8), ("Pass 2 distance (m)", 9), ("Pass 2 misalignment", 9), ("Text direction found", 9)]
    header(ws, 1, [c for c, _ in LCOLS], [w for _, w in LCOLS])
    lrows = []
    for li, l in enumerate(labels):
        want = wanted(l)
        p1 = m1.get(li)
        p2 = m2.get(li)
        o1 = objs_by_index(p1)
        lrows.append({"Label": li + 1, "Text": l["raw"], "Kind": l["kind"], "Names class": " / ".join(sorted(want)),
                      "x (px)": round(l["c"].x * s, 1), "y (px)": round(l["c"].y * s, 1),
                      "Pass 1 object": o1["id"] if o1 else "unmatched",
                      "Pass 1 cluster": o1["cluster"] if o1 else "",
                      "Pass 1 cluster class": (o1["cls"] or "unnamed") if o1 else "",
                      "Pass 1 agrees": ("yes" if o1["cls"] in want else ("unnamed" if o1["cls"] is None else "no")) if o1 else "",
                      "Pass 2 object": objs_by_index(p2)["id"] if p2 else "unmatched",
                      "Pass 2 distance (m)": p2[2][0] if p2 else None, "Pass 2 misalignment": p2[2][1] if p2 else None,
                      "Text direction found": "yes" if l["dir"] is not None else "no"})
    for r, row in enumerate(lrows, 2):
        for ci, (name, _) in enumerate(LCOLS, 1):
            ws.cell(row=r, column=ci, value=row[name]).font = body
    NL = len(lrows) + 1
    LE = lambda name: f"Labels!${cletter(LCOLS, name)}$2:${cletter(LCOLS, name)}${NL}"
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(LCOLS))}{NL}"

    # --- Clusters ---
    ws = wb.create_sheet("Clusters", 1)
    kinds = sorted({l["kind"] for l in labels})
    CC = ["Cluster", "Drawing kind", "Members", "Length min (cm)", "Length max (cm)", "Width min (cm)", "Width max (cm)",
          "Touch solids", "Inside an outline", "In a strip band", "On a stamped region"] + \
         [f"Votes: {kd}" for kd in kinds] + ["Votes total", "Top share", "Class", "Named by", "Silhouette"]
    header(ws, 1, CC, [8, 16, 8, 9, 9, 9, 9, 9, 9, 9, 9] + [9] * len(kinds) + [8, 8, 17, 30, 9])
    for r, c in enumerate(info, 2):
        ms = [o for o in objs if o["cluster"] == c["id"]]
        cl = lambda n: cletter(OBJ_COLS, n)
        crit = f'{OE("Cluster")},$A{r}'
        vals = [c["id"], c["kind"], f"=COUNTIFS({crit})",
                f'=_xlfn.MINIFS({OE("Length (cm)")},{crit})', f'=_xlfn.MAXIFS({OE("Length (cm)")},{crit})',
                f'=_xlfn.MINIFS({OE("Width (cm)")},{crit})', f'=_xlfn.MAXIFS({OE("Width (cm)")},{crit})',
                f'=IFERROR(COUNTIFS({crit},{OE("Touches solids")},">0")/$C{r},0)',
                f'=IFERROR(COUNTIFS({crit},{OE("Inside an outline")},1)/$C{r},0)',
                f'=IFERROR(COUNTIFS({crit},{OE("In a strip band")},1)/$C{r},0)',
                f'=IFERROR(COUNTIFS({crit},{OE("On a stamped region")},1)/$C{r},0)']
        # votes counted from the label sheet, compound labels speaking for two objects included
        for kd in kinds:
            vals.append(c["votes"].get(kd, 0) + (c["votes"].get("strip_footing", 0) if False else 0))
        vstart, vend = get_column_letter(12), get_column_letter(11 + len(kinds))
        vals += [f"=SUM({vstart}{r}:{vend}{r})", f"=IFERROR(MAX({vstart}{r}:{vend}{r})/{get_column_letter(12 + len(kinds))}{r},0)",
                 c["name"] or "unnamed", c["named_by"] or "no label votes", c["silhouette"]]
        for ci, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=ci, value=v)
            cell.font = body
        for ci in (8, 9, 10, 11, 13 + len(kinds)):
            ws.cell(row=r, column=ci).number_format = "0%"
    nr = len(info) + 3
    ws.cell(row=nr, column=1, value="Votes are the first, neutral pass (program output). A 'VP..-SF..' label votes for the strip it "
                                    "sits on and for the wall standing in that strip.").font = note
    ws.freeze_panes = "C2"

    # --- Summary by building ---
    ws = wb.create_sheet("Summary by building", 1)
    SC = ["Building", "Class", "Type key", "Count", "Length min (cm)", "Length max (cm)", "Width min (cm)",
          "Width max (cm)", "Same section?", "Total length (m)", "Of which typed by inference"]
    header(ws, 1, SC, [9, 17, 16, 7, 9, 9, 9, 9, 13, 10, 11])
    main = [o for o in objs if o["cls"]]
    combos = sorted({(str(o["building"] or "unknown"), o["cls"], o["key"] or "(none)") for o in main},
                    key=lambda t: (t[0], order.get(t[1], 99), t[2]))
    for r, (b, cls, key) in enumerate(combos, 2):
        crit = f'{OE("Building")},$A{r},{OE("Class")},$B{r},{OE("Type key")},$C{r}'
        vals = [b, cls, key, f"=COUNTIFS({crit})",
                f'=_xlfn.MINIFS({OE("Length (cm)")},{crit})', f'=_xlfn.MAXIFS({OE("Length (cm)")},{crit})',
                f'=_xlfn.MINIFS({OE("Width (cm)")},{crit})', f'=_xlfn.MAXIFS({OE("Width (cm)")},{crit})',
                (f'=IF(OR(C{r}="(none)",B{r}="raft",B{r}="sump_pit"),"n/a",IF(OR(B{r}="column",B{r}="shear_wall",'
                 f'B{r}="isolated_footing",B{r}="combined_footing"),IF(AND(F{r}-E{r}<=3,H{r}-G{r}<=3),"yes","no"),'
                 f'IF(H{r}-G{r}<=3,"yes (thickness)","no (thickness)")))'),
                f'=SUMIFS({OE("Length (cm)")},{crit})/100',
                f'=COUNTIFS({crit},{OE("Type source")},"inferred*")']
        for ci, v in enumerate(vals, 1):
            ws.cell(row=r, column=ci, value=v).font = body
        ws.cell(row=r, column=10).number_format = "0.0"
    last = len(combos) + 1
    ws.cell(row=last + 2, column=1, value="All").font = bold
    ws.cell(row=last + 2, column=4, value=f"=SUM(D2:D{last})").font = bold
    ws.cell(row=last + 3, column=1, value="Unnamed clusters are not listed here (see 'Clusters'). 'Same section?' allows 3 cm; "
                                          "columns, shear walls and footings compare length and width, walls and bands of "
                                          "varying length compare thickness only.").font = note
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:K{last}"

    # --- Conflicts ---
    ws = wb.create_sheet("Conflicts")
    CF = ["Label", "Text", "Names class", "x (px)", "y (px)", "Issue", "Object it sits by (pass 1)",
          "That object's class", "Pass 2 result"]
    header(ws, 1, CF, [6, 20, 24, 8, 8, 44, 10, 17, 16])
    r = 2
    for row in lrows:
        issue = None
        if row["Pass 1 agrees"] == "no":
            issue = "first, neutral pass put it on an object of another class"
        if row["Pass 2 object"] == "unmatched":
            issue = (issue + "; " if issue else "") + "no object of its class within reach in the second pass"
        if issue:
            for ci, v in enumerate([row["Label"], row["Text"], row["Names class"], row["x (px)"], row["y (px)"], issue,
                                    row["Pass 1 object"], row["Pass 1 cluster class"], row["Pass 2 object"]], 1):
                ws.cell(row=r, column=ci, value=v).font = body
            r += 1
    ws.freeze_panes = "C2"

    # --- Validation ---
    ws = wb.create_sheet("Validation")
    ws["A1"] = "Geometric classes against the verified (colour-based) detections"
    ws["A1"].font = title
    classes = ["column", "shear_wall", "retaining_wall", "isolated_footing", "combined_footing", "crane_footing",
               "mass_concrete_pad"]
    got = ["column", "shear_wall", "retaining_wall", "isolated_footing", "combined_footing", "strip_footing",
           "grade_beam", "raft", "sump_pit", "unnamed"]
    header(ws, 3, ["Verified class"] + got + ["Not found", "Total", "Correct"], [18] + [11] * (len(got) + 3))
    for r, vc in enumerate(classes, 4):
        ws.cell(row=r, column=1, value=vc).font = body
        for ci, g in enumerate(got, 2):
            ws.cell(row=r, column=ci, value=f'=COUNTIFS({OE("(validation) verified class")},$A{r},{OE("Class")},'
                                            f'{get_column_letter(ci)}$3)').font = body
        nf = sum(1 for v, c in ver if v == vc and c == "not found")
        ws.cell(row=r, column=len(got) + 2, value=nf).font = body
        ws.cell(row=r, column=len(got) + 3, value=f"=SUM(B{r}:{get_column_letter(len(got) + 2)}{r})").font = bold
        want = VERIFIED_TO_CLASS.get(vc)
        if want:
            col = get_column_letter(got.index(want) + 2)
            c = ws.cell(row=r, column=len(got) + 4, value=f"=IFERROR({col}{r}/{get_column_letter(len(got) + 3)}{r},0)")
            c.number_format, c.font = "0.0%", bold
    rr = len(classes) + 5
    ws.cell(row=rr, column=1, value="Crane footings count as isolated footings. Concrete pads carry no label, so no class "
                                    "can be expected for them. 'Not found' = no object overlapping the verified shape "
                                    "by 80%.").font = note
    ws.cell(row=rr + 2, column=1, value="Band widths found in the drawing").font = bold
    header(ws, rr + 3, ["Width (m)", "Paired line length (m)", "Share filled", "Cluster class"], [18, 16, 12, 18])
    for i, bf in enumerate(log["band_families"], rr + 4):
        cl = next((c["name"] or "unnamed" for c in info if c["kind"] == f"band {bf['width']:.2f} m"), "")
        for ci, v in enumerate([bf["width"], bf["paired_m"], bf["filled_share"], cl], 1):
            ws.cell(row=i, column=ci, value=v).font = body
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)
    return lrows


_OBJ_INDEX = []


def objs_by_index(p):
    return _OBJ_INDEX[p[0]] if p else None


CLASS_NAME = {"column": "Column", "shear_wall": "Shear wall", "retaining_wall": "Retaining wall", "core_wall": "Core wall",
              "isolated_footing": "Isolated footing", "combined_footing": "Combined footing", "strip_footing": "Strip footing",
              "grade_beam": "Grade beam", "raft": "Raft", "sump_pit": "Sump pit"}


def display_label(o):
    """The text shown with an element: its class in words, its type (~ = inferred) and its measured
    size in metres (length x width of its box; a shape that is not a rectangle also gets its area)."""
    if not o["cls"]:
        return f"C{o['cluster']}"
    key = o["key"] or ""
    if key and o["type_source"].startswith("inferred"):
        key = "~" + key
    f = o["feat"]
    size = f"{f['length']:.2f} × {f['width']:.2f} m"
    if not box_fits(o):
        size += f", {f['area']:.1f} m²"
    return " ".join(x for x in (CLASS_NAME.get(o["cls"], o["cls"]), key, size) if x)


# the brief's element categories; the class is kept as the subclass
ELEMENT_TYPE = {"column": "column", "shear_wall": "wall", "retaining_wall": "wall", "core_wall": "wall",
                "isolated_footing": "footing", "combined_footing": "footing", "strip_footing": "footing",
                "grade_beam": "beam", "raft": "slab_raft", "sump_pit": "other"}


def write_json(path, pdf, objs, calib, dpi):
    """Machine-readable results: every classified object with its class, type and boxes in pixels
    of the page rendered at `dpi` (origin top-left, y down). Unnamed objects are only counted."""
    s = dpi / 72.0
    px = lambda pts: [[round(x * s, 1), round(y * s, 1)] for x, y in pts]
    page = pymupdf.open(pdf)[0]
    dets = []
    for o in sorted((o for o in objs if o["cls"]), key=lambda o: o["id"]):
        f = o["feat"]
        x0, y0, x1, y1 = o["g"].bounds
        d = {"id": o["id"], "element_type": ELEMENT_TYPE.get(o["cls"], "other"), "class": o["cls"],
             "display_label": display_label(o),
             "type_key": ("~" if o["type_source"].startswith("inferred") else "") + (o["key"] or "?"),
             "type_source": o["type_source"], "label": o["label_text"] or None,
             "class_named_by": o.get("wall_rule") or o.get("named_by") or "",
             "building": o["building"], "drawing_kind": o["kind"],
             "bbox_px": [round(x0 * s, 1), round(y0 * s, 1), round(x1 * s, 1), round(y1 * s, 1)],
             "obb_px": px(o["o"]["obb"]) if o["o"] else None,
             "size_m": {"length": round(f["length"], 3), "width": round(f["width"], 3)},
             "angle_deg": round(o["o"]["angle_deg"], 1) if o["o"] else None, "area_m2": round(f["area"], 3)}
        if not box_fits(o):
            d["polygon_px"] = px(o["g"].exterior.coords)
        if o.get("net") is not None:
            d["net_length_m"] = round(o["net"], 2)
        dets.append(d)
    out = {"source_pdf": Path(pdf).name,
           "image": {"dpi": dpi, "width_px": round(page.rect.width * s), "height_px": round(page.rect.height * s),
                     "coordinates": "pixels of the page rendered at this dpi, origin top-left, y down"},
           "calibration": {"pt_per_m": round(calib["pt_per_m"], 4), "scale": f"1/{round(72 / 0.0254 / calib['pt_per_m'])}"},
           "counts": dict(collections.Counter(d["class"] for d in dets)),
           "unnamed_objects": sum(1 for o in objs if not o["cls"]),
           "detections": dets}
    Path(path).write_text(json.dumps(out, indent=1, ensure_ascii=False))


def box_fits(o):
    """Draw an object as its oriented box? Shapes traced from the drawing's lines (outlines, rafts)
    are drawn as traced unless they are rectangles (a small notch still shows). Objects built from fills, hatch strokes or line
    pairs are straight pieces by construction (bent ones were split into legs); their box is drawn
    if the shape stays within one thickness of it (a hatched wall's ragged stroke ends)."""
    if not o["o"]:
        return False
    if o["kind"] in ("outline", "stamped_region"):
        return o["o"]["rectness"] >= 0.99            # a rectangle within drafting precision
    obb = o["o"]["obb"]
    thickness = min(math.dist(obb[0], obb[1]), math.dist(obb[1], obb[2]))
    return o["g"].exterior.hausdorff_distance(LineString(list(obb) + [obb[0]])) <= thickness


def label_anchor(o, gap=0.6):
    """Where a label is written: along a long element, just outside the long side higher on the
    page, read left to right; above the box's corner for a compact element."""
    oo, f = o["o"], o["feat"]
    if oo and f["width"] and f["length"] >= 3 * f["width"] and box_fits(o):
        c = oo["obb"]
        sides = [(c[i], c[(i + 1) % 4]) for i in range(4)]
        longs = sorted(sides, key=lambda sd: -math.dist(*sd))[:2]
        top = min(longs, key=lambda sd: sd[0][1] + sd[1][1])
        a, b = sorted(top, key=lambda q: (q[0], -q[1]))
        ang = math.degrees(math.atan2(b[1] - a[1], b[0] - a[0]))
        nx, ny = math.sin(math.radians(ang)), -math.cos(math.radians(ang))     # the side's outward normal, up the page
        return (a[0] + gap * nx, a[1] + gap * ny), ang
    x0, y0, _, _ = o["g"].bounds
    return (x0, y0 - gap), 0.0


def write_pdf(pdf, out_path, objs):
    doc = pymupdf.open(pdf)
    page = doc[0]
    layers = {}
    for o in objs:
        name = o["cls"] or "unnamed"
        if name not in layers:
            layers[name] = doc.add_ocg(f"GEO {name}", on=name != "unnamed")
        col = CLASS_COLOURS.get(o["cls"], UNNAMED_COLOUR)
        if box_fits(o):
            pts = [pymupdf.Point(x, y) for x, y in o["o"]["obb"]]
            page.draw_polyline(pts + pts[:1], color=col, width=0.5, oc=layers[name])
        else:   # not a rectangle (rafts, hatched regions): its real outline, and its box dashed
            pts = [pymupdf.Point(x, y) for x, y in o["g"].exterior.coords]
            page.draw_polyline(pts, color=col, width=0.6, oc=layers[name])
            x0, y0, x1, y1 = o["g"].bounds
            page.draw_rect(pymupdf.Rect(x0, y0, x1, y1), color=col, width=0.25, dashes="[2 2] 0", oc=layers[name])
        (ax, ay), ang = label_anchor(o)
        pt = pymupdf.Point(ax, ay)
        page.insert_text(pt, display_label(o), fontsize=2.6, color=col, oc=layers[name],
                         morph=(pt, pymupdf.Matrix(-ang)) if ang else None)   # page y points down
    doc.save(out_path, garbage=3, deflate=True)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("pdf")
    ap.add_argument("--config", default=str(Path(__file__).with_name("plan_config.toml")))
    ap.add_argument("--reference", default="out/detections.json")
    ap.add_argument("--out", default="out")
    ap.add_argument("--dpi", type=int, default=150)
    ap.add_argument("--version", type=int, default=None, help="output version number (default: next free one)")
    args = ap.parse_args()
    cfg = tomllib.loads(Path(args.config).read_text())
    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)
    r = run(args.pdf, cfg, args.reference)
    objs, info, labels, m1, m2, log, ver, calib = (r[n] for n in ("objs", "info", "labels", "m1", "m2", "log", "ver", "calib"))
    _OBJ_INDEX[:] = objs
    # a new version number per run, so earlier results stay for comparison
    done = [int(m.group(1)) for f in out.glob("geometry_classes_v*.pdf") if (m := re.search(r"_v(\d+)\.pdf$", f.name))]
    v = args.version or (max(done) + 1 if done else 1)
    lrows = write_workbook(out / f"geometry_classes_v{v}.xlsx", list(objs), info, labels, m1, m2, log, ver, args.dpi, calib)
    write_pdf(args.pdf, out / f"geometry_classes_v{v}.pdf", objs)
    write_json(out / f"geometry_classes_v{v}.json", args.pdf, objs, calib, args.dpi)
    pymupdf.open(out / f"geometry_classes_v{v}.pdf")[0].get_pixmap(dpi=args.dpi).save(out / f"geometry_classes_v{v}.png")
    print(f"written out/geometry_classes_v{v}.xlsx, .pdf, .json and .png")
    cls = collections.Counter(o["cls"] or "unnamed" for o in objs)
    print("objects:", len(objs), dict(cls))
    print("clusters:", len(info), "named:", sum(1 for c in info if c["name"]))
    print("joints:", log["joints"], "regions:", log["regions"])
    agree = collections.Counter(r["Pass 1 agrees"] for r in lrows)
    print("pass 1 labels:", dict(agree), " pass 2 unmatched:", sum(1 for r in lrows if r["Pass 2 object"] == "unmatched"))
    vc = collections.Counter((v, c) for v, c in ver)
    for v in ("column", "shear_wall", "retaining_wall", "isolated_footing", "combined_footing", "crane_footing", "mass_concrete_pad"):
        print("verified", v, "->", {c: n for (vv, c), n in vc.items() if vv == v})
    return 0


def run(pdf, cfg, reference=None):
    """The whole detection and classification, without writing anything."""
    doc, page, paths, texts = ds.load_page(pdf)
    calib = ds.calibrate(paths, texts, ds.glyph_size(texts, cfg), cfg)
    k = calib["pt_per_m"]
    objs, segs, tiny, log = objects(paths, texts, k)
    features(objs, tiny, k, circles(paths, k))
    info = cluster(objs)
    labels = read_labels(texts, paths, cfg, k, es.text_index(texts))
    m1 = match(labels, objs, k)
    name_clusters(info, labels, objs, m1, label_size_meaning(cfg))

    # Column-shaped pieces drawn like walls are set aside before labels are attached: how columns are
    # drawn is learned from what the P labels land on in the neutral pass (the pen of their outline);
    # a column-cluster member with an outline pen none of those have, or that continues a wall of the
    # same thickness, is a wall piece, not a column (no label needed for this decision).
    col_pens = collections.Counter(objs[i]["outline_pen"] for li, (i, cost, _) in m1.items()
                                   if labels[li]["kind"] == "column" and objs[i]["cls"] == "column")
    usual_pens = {p for p, n in col_pens.items() if n >= G.pen_share * sum(col_pens.values())}
    solids_all = [o for o in objs if o["kind"] == "solid"]
    for o in objs:
        if o["cls"] != "column":
            continue
        continues_wall = any(q is not o and q["cls"] != "column" and q["g"].distance(o["g"]) < G.touch_m * k
                             and abs(q["feat"]["width"] - o["feat"]["width"]) <= G.width_tol for q in solids_all)
        if (usual_pens and o["outline_pen"] not in usual_pens) or continues_wall:
            o["cls"], o["named_by"] = "core_wall", "rule: column-shaped piece drawn like a wall (outline pen or continuity)"
            o["piece"] = True

    voted = [c for c in info if c["named_by"] == "label vote" and c["centroid"] is not None]

    def nearest_voted_class(o):
        cands = [c for c in voted if c["kind"] == o["kind"]]
        if "x" not in o or not cands:
            return None
        return min(cands, key=lambda c: float(np.linalg.norm(c["centroid"] - o["x"])))["name"]

    def compatible(l, o):
        if not dims_agree(l, o):
            return None
        # a label is attached only to objects of its home drawing kind; a 'VP..-SF..' label names a
        # strip (band) and the wall in it (solid)
        home = HOME.get(l["kind"], set()) | ({"band", "solid"} if "parts" in l else set())
        if home and o["kind"] not in home:
            return None
        if o["cls"] in wanted(l):
            return 0.0
        # an unnamed cluster, or a class only guessed, is open to a label at a cost - if the object is
        # geometrically closest to a cluster that labels named with that label's class
        weak = o["cls"] is None or not o.get("named_by", "").startswith("label vote")
        if not weak:
            return None
        near = nearest_voted_class(o)
        return G.unnamed_penalty if near is None or FAMILY.get(near) in {FAMILY.get(c) for c in wanted(l)} else None

    # second pass in two steps: first within classes, which fixes the size of each type from the
    # labels sitting right on their element; then again, penalising an object whose size is not its
    # label's type size, or that lies in a joint region of another building than the label's
    m2a = match(labels, objs, k, compatible)
    sizes = collections.defaultdict(list)
    for li, (i, cost, (d, mis)) in m2a.items():
        l = labels[li]
        if l["kind"] in ("column", "wall", "footing:S", "footing:SC") and d <= G.sure_distance:
            sizes[l["key"]].append((d, objs[i]["feat"]["length"], objs[i]["feat"]["width"]))
    # when the close matches of a type disagree, the closest one decides, with those agreeing with it
    type_size = {}
    for key, vs in sizes.items():
        vs.sort()
        _, L0, W0 = vs[0]
        agree = [(L, W) for _, L, W in vs if abs(L - L0) <= G.size_tol and abs(W - W0) <= G.size_tol]
        type_size[key] = (statistics.median(a[0] for a in agree), statistics.median(a[1] for a in agree))
    regions, n_joints = building_regions(objs, labels, segs, k)
    log["joints"], log["regions"] = n_joints, [(r["building"], r["votes"]) for r in regions]
    for o in objs:
        c = o["g"].centroid
        reg = next((r for r in regions if r["g"].contains(c)), None)
        # at a joint both buildings' elements stand side by side: only objects well inside a region
        # are told its building
        inside = reg is not None and reg["g"].exterior.distance(c) >= G.joint_margin * k
        o["region_bld"] = reg["building"] if inside else None

    def checked(l, o):
        pen = compatible(l, o)
        if pen is None:
            return None
        if l["key"] in type_size:
            L, W = type_size[l["key"]]
            if abs(o["feat"]["length"] - L) > G.size_tol or abs(o["feat"]["width"] - W) > G.size_tol:
                pen += G.size_penalty
        if l.get("building") and o["region_bld"] and o["region_bld"] != l["building"]:
            pen += G.region_penalty
        return pen

    m2 = match(labels, objs, k, checked)
    assign_types(objs, labels, m2, k, cfg, regions)
    # evidence ranks: a cluster's own label vote > the element's own label > a similarity guess
    for o in objs:
        src = o.get("label_src")
        if not src or not o["type_source"].startswith("printed"):
            continue
        lab_cls = LABEL_CLASS.get(src["kind"])
        if lab_cls and lab_cls != o["cls"] and not o.get("named_by", "").startswith("label vote"):
            o["cls"] = lab_cls
            o["named_by"] = "its own label (the cluster's class was only a guess)"
    # Unlabelled members, final rules (no label needed, geometry decides):
    # - a column-cluster member drawn unlike the labelled columns (an outline pen none of them has)
    #   or continuing a wall of the same thickness is a wall piece, not a column;
    # - an unlabelled wall-like solid or hatched wall along the building edge is a retaining wall,
    #   one inside the building that is no V type is a core wall.
    wall_w = statistics.median(o["feat"]["width"] for o in objs if o["cls"] in ("shear_wall", "retaining_wall")) \
        if any(o["cls"] in ("shear_wall", "retaining_wall") for o in objs) else G.wall_default
    # a plan without P labels cannot name its columns by vote: there, an unlabelled solid of column
    # shape (EN 1992-1-1 5.3.1(7): section no longer than 4 x its width) is a column, unless it
    # continues a wall of the same thickness (a piece of that wall)
    p_columns = any(c["name"] == "column" and c["named_by"] == "label vote" for c in info)
    solids_all = [o for o in objs if o["kind"] == "solid"]
    for o in objs:
        f = o["feat"]
        if p_columns or o["kind"] != "solid" or not f["width"] or f["length"] > es.T.ec2_ratio * f["width"] \
                or o["type_source"].startswith("printed") or o.get("named_by", "").startswith("label vote"):
            continue
        if any(q is not o and q["g"].distance(o["g"]) < G.touch_m * k and abs(q["feat"]["width"] - f["width"]) <= G.width_tol
               and q["feat"]["length"] > es.T.ec2_ratio * q["feat"]["width"] for q in solids_all):
            o["piece"] = True
            continue
        o["cls"], o["wall_rule"] = "column", "rule: column by shape (EN 1992-1-1, L <= 4 W); the plan has no P labels to name columns"
    for o in objs:
        if o["cls"] == "column" and not o.get("piece"):
            continue
        f = o["feat"]
        # a wall drawn only by its outline, with a pattern inside (dots or strokes): wall-shaped
        # (longer than 4 x its width, EN 1992) and no thicker than the plan's walls
        outline_wall = o["kind"] == "outline" and f["width"] and f["length"] > es.T.ec2_ratio * f["width"] \
            and f["width"] <= 2 * wall_w and f["dots_inside"] >= 1
        wall_like = o["kind"] == "cross_hatch" or o.get("piece") or outline_wall or \
            (o["kind"] == "solid" and f["width"] <= 2 * wall_w and f["length"] >= 2 * f["width"])
        if not wall_like or o["type_source"].startswith("printed"):
            continue
        if o.get("piece") or (o["cls"] in (None, "shear_wall", "core_wall", "unconfirmed") and not (o["cls"] == "shear_wall" and o["key"])):
            # along the building edge, or standing on a strip footing (strips run under edge walls)
            if o["feat"]["on_perimeter"] or o["feat"]["in_band_holding_solid"]:
                o["cls"], o["wall_rule"] = "retaining_wall", "rule: unlabelled wall along the building edge or on a strip footing"
            else:
                o["cls"] = "core_wall"
                o["wall_rule"] = o.get("named_by") if o.get("piece") else "rule: unlabelled interior wall, no V type of its size"
            if outline_wall:
                o["wall_rule"] += " (drawn as an outline with a pattern inside)"
    # retaining walls found by the edge rule take the type of their thickness, learned from the walls
    # whose printed labels name a type of that class on this plan (VP20 = the 20 cm ones here)
    thick = collections.defaultdict(list)
    for o in objs:
        if o["cls"] == "retaining_wall" and o["key"] and o["type_source"].startswith("printed"):
            thick[o["key"]].append(o["feat"]["width"])
    for o in objs:
        if o["cls"] == "retaining_wall" and not o["key"] and o.get("wall_rule"):
            fits = sorted((abs(statistics.median(ws) - o["feat"]["width"]), key) for key, ws in thick.items()
                          if abs(statistics.median(ws) - o["feat"]["width"]) <= G.width_tol)
            if fits:
                o["key"] = "|".join(key for _, key in fits)
                o["type_source"] = "inferred: the thickness of this retaining-wall type" + (" (several fit, nearest first)" if len(fits) > 1 else "")
    ver = validate(objs, reference, k) if reference else None
    return {"objs": objs, "info": info, "labels": labels, "m1": m1, "m2": m2, "log": log, "ver": ver,
            "calib": calib, "regions": regions, "k": k}


if __name__ == "__main__":
    sys.exit(main())
