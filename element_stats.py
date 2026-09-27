#!/usr/bin/env python3
"""Colour- and layer-free catalogue of the structural elements on a foundation plan.

Question this answers: can each element be recognised from geometry alone (shape, size,
drawing structure, relative position, text), with no colour and no CAD layer? It detects
every element family that way, measures the colour-free marks of each element, checks the
result against the verified detections of detect_structures.py (run that first: it writes
out/detections.json), and writes a per-building catalogue with counts and dimensions.

Colour-free cues used
    paint order     AutoCAD writes each object's pieces one after another, so consecutive
                    filled pieces that touch are one drawn object (splits touching elements)
    shape           rectangle, size, aspect ratio (EN 1992-1-1 5.3.1(7): a section whose long
                    side is at most 4x its short side is a column, otherwise a wall)
    drawing type    solid fill / cross-hatch / closed outline / pair of parallel lines / dash chain
    text            label grammar (P, V, VP, S, SC, SF, LG, Radier), sizes written in labels
    relations       inside a footing, on a raft, touching walls, at grade-beam junctions, on grid lines

Colour and layer are read only in the validation columns, never to detect or classify.

Usage
    python element_stats.py "GHF_EXE_PLN_STR (1)-FONDATIONS (1).pdf" --out out
"""
import argparse
import collections
import json
import math
import re
import statistics
import sys
import tomllib
from pathlib import Path
from types import SimpleNamespace

from shapely.geometry import LineString, MultiPoint, Point, Polygon, box
from shapely.ops import unary_union
from shapely.strtree import STRtree

import detect_structures as ds

# ---------------------------------------------------------------------------
# Tolerances (metres unless stated). Kept here because they describe drawing
# practice in general, not this plan.
# ---------------------------------------------------------------------------
T = SimpleNamespace(
    touch=0.02,                 # pieces closer than this touch
    paint_gap=3,                # pieces at most this many paths apart in paint order can be one object
    min_side=0.15,              # smallest structural section side
    rectness=0.95,              # area / minimum rotated rectangle area
    dup_iou=0.9,                # objects overlapping more than this are one object drawn twice
    ec2_ratio=4.0,              # EN 1992-1-1 5.3.1(7) column/wall limit
    hatch_len=(0.05, 1.0),      # hatch stroke length range
    hatch_neighbours=3,         # a hatch stroke has at least this many parallel twins nearby
    hatch_min_strokes=20,
    band_tol=0.01,              # band width tolerance
    band_min_overlap=0.5,
    band_run_gap=0.6,           # band pieces closer than this along their axis are one run
    band_empty=0.3,             # a beam band is at most this share covered by filled objects
    band_wall_share=0.12,       # a strip-footing band has walls covering at least this share of it
    grid_min_pieces=8,
    grid_min_extent=8.0,
    footing_min_area=0.5,       # m2
    footing_max_aspect=3.0,
    pad_side=(0.4, 1.2),
    pad_min_dots=5,
    label_radius=1.5,
    building_k=5,
    junction_radius=0.5,
)


# ---------------------------------------------------------------------------
# Primitives
# ---------------------------------------------------------------------------
def text_index(texts):
    boxes = [box(*t["rect"]).buffer(1.0) for t in texts]
    return boxes, STRtree(boxes)


def in_text(geom, tboxes):
    """A stroke is part of a text if it lies inside a text's box and is no longer than the text
    is tall (a label written between two lines must not swallow the lines)."""
    boxes, tree = tboxes
    for i in tree.query(geom):
        b = boxes[i]
        if b.contains(geom):
            x0, y0, x1, y1 = b.bounds
            if max(geom.bounds[2] - geom.bounds[0], geom.bounds[3] - geom.bounds[1]) <= 1.2 * min(x1 - x0, y1 - y0):
                return True
    return False


def is_closed(path):
    """True if the path draws a closed shape: a rectangle/quad item, or a chain of segments that
    ends where it starts (an open polyline is not closed, however many segments it has)."""
    chain = []
    for kind, pts in path["items"]:
        if kind in ("re", "qu"):
            return True
        if chain and math.dist(chain[-1], pts[0]) < 1e-3:
            chain += pts[1:]
        else:
            if len(chain) >= 4 and math.dist(chain[0], chain[-1]) < 1e-3:
                return True
            chain = list(pts)
    return len(chain) >= 4 and math.dist(chain[0], chain[-1]) < 1e-3


def straight_segments(paths, k, tboxes, min_m=0.0, max_m=math.inf):
    """Every straight stroke segment, with its paint index; text strokes are skipped."""
    out = []
    for i, p in enumerate(paths):
        if p["type"] not in ("s", "fs"):
            continue
        closed = is_closed(p)                   # the stroke is a side of a closed shape
        for ls in ds.path_lines(p):
            L = ls.length / k
            if min_m <= L <= max_m and not in_text(ls, tboxes):
                (ax, ay), (bx, by) = ls.coords[0], ls.coords[-1]
                out.append({"g": ls, "a": (ax, ay), "b": (bx, by), "len": L, "idx": i, "stroke": p["type"] == "s",
                            "closed": closed,
                            "ang": math.atan2(by - ay, bx - ax) % math.pi, "w": round(p["width"], 2)})
    return out


def axis(ang):
    return math.cos(ang), math.sin(ang)


def along(pt, ang):
    ux, uy = axis(ang)
    return pt[0] * ux + pt[1] * uy


def offset(pt, ang):
    ux, uy = axis(ang)
    return -pt[0] * uy + pt[1] * ux


def mark_dashed(segs, k):
    """Flag segments that are pieces of a dashed line (many collinear pieces with gaps). Dashed
    lines are axes or hidden lines, never the visible edge of an element."""
    groups = collections.defaultdict(list)
    for i, s_ in enumerate(segs):
        if s_["len"] > 2.0:
            continue
        a = round(math.degrees(s_["ang"]) * 5) / 5 % 180
        mid = ((s_["a"][0] + s_["b"][0]) / 2, (s_["a"][1] + s_["b"][1]) / 2)
        groups[(a, round(offset(mid, s_["ang"]) / (0.03 * k)))].append(i)
    n = 0
    for ids in groups.values():
        if len(ids) < T.grid_min_pieces:
            continue
        ang = segs[ids[0]]["ang"]
        iv = sorted((min(along(segs[i]["a"], ang), along(segs[i]["b"], ang)),
                     max(along(segs[i]["a"], ang), along(segs[i]["b"], ang))) for i in ids)
        lo, hi = iv[0][0], max(v[1] for v in iv)
        cover = sum(v[1] - v[0] for v in iv) / (hi - lo) if hi > lo else 1
        if cover <= 0.97:
            for i in ids:
                segs[i]["dashed"] = True
                n += 1
    # the pen the dashed axes are drawn with; continuous lines of that pen are axes too
    pen = collections.Counter(s_["w"] for s_ in segs if s_.get("dashed") and s_["w"] > 0).most_common(1)
    pen = pen[0][0] if pen else None
    for s_ in segs:
        if pen is not None and s_["w"] == pen:
            s_["axis_pen"] = True
    return n, pen


# ---------------------------------------------------------------------------
# 1. Filled objects, grouped by paint order
# ---------------------------------------------------------------------------
def filled_objects(paths, k):
    objs, cur, last = [], None, -10
    for i, p in enumerate(paths):
        if p["type"] not in ("f", "fs") or not p["fill"]:
            continue
        parts = ds.subpolygons(p)
        if not parts:
            continue
        g = unary_union(parts)
        if cur is not None and i - last <= T.paint_gap and cur["g"].distance(g) < T.touch * k:
            cur["g"] = cur["g"].union(g)
            cur["paths"].append(i)
        else:
            cur = {"g": g, "paths": [i]}
            objs.append(cur)
        last = i
    rects, others = [], []
    for o in objs:
        g = o["g"]
        if g.geom_type != "Polygon":
            g = g.buffer(0.01 * k).buffer(-0.01 * k)
        if g.geom_type != "Polygon" or g.is_empty:
            if not g.is_empty:
                others.append(g)
            continue
        oo = ds.oriented(g, k)
        if oo["rectness"] > T.rectness and oo["width_m"] >= T.min_side:
            rects.append({"g": g, "o": oo, "paths": o["paths"]})
        else:
            others.append(g)
    kept, dups = [], 0
    for r in sorted(rects, key=lambda r: -r["g"].area):
        if any(r["g"].intersection(q["g"]).area / r["g"].union(q["g"]).area > T.dup_iou for q in kept):
            dups += 1
            continue
        kept.append(r)
    return kept, dups, others


def straight_stretches(g, k, t_range=(0.1, 0.6), min_len=0.3):
    """Split a thin filled polygon with bends (a wall drawn as one object around corners) into
    straight stretches: pairs of opposite, parallel edges one wall-thickness apart."""
    ring = list(g.simplify(0.01 * k).exterior.coords)
    edges = []
    for a, b in zip(ring, ring[1:]):
        L = math.dist(a, b)
        if L / k >= min_len:
            edges.append((a, b, math.atan2(b[1] - a[1], b[0] - a[0]) % math.pi))
    out = []
    for i, (a, b, ang) in enumerate(edges):
        for c, d, ang2 in edges[i + 1:]:
            if abs(math.sin(ang2 - ang)) > 0.02:
                continue
            t = abs(offset(((c[0] + d[0]) / 2, (c[1] + d[1]) / 2), ang) - offset(a, ang)) / k
            if not t_range[0] <= t <= t_range[1]:
                continue
            lo = max(min(along(a, ang), along(b, ang)), min(along(c, ang), along(d, ang)))
            hi = min(max(along(a, ang), along(b, ang)), max(along(c, ang), along(d, ang)))
            if (hi - lo) / k < min_len:
                continue
            ux, uy = axis(ang)
            o1, o2 = offset(a, ang), offset(((c[0] + d[0]) / 2, (c[1] + d[1]) / 2), ang)
            P = lambda u, v: (u * ux - v * uy, u * uy + v * ux)
            piece = Polygon([P(lo, o1), P(hi, o1), P(hi, o2), P(lo, o2)])
            if piece.is_valid and g.buffer(0.01 * k).contains(piece):   # both faces of the same stretch
                out.append(piece)
    kept = []
    for p_ in sorted(out, key=lambda p_: -p_.area):
        if all(p_.intersection(q).area < 0.5 * p_.area for q in kept):
            kept.append(p_)
    return kept


def ec2_class(o):
    """Nominal section (rounded to the cm) against the EN 1992-1-1 column/wall limit."""
    L, W = round(o["length_m"] * 100), round(o["width_m"] * 100)
    return "column" if L <= T.ec2_ratio * W else "wall"


# ---------------------------------------------------------------------------
# 2. Cross-hatched regions (walls drawn with a hatch instead of a fill)
# ---------------------------------------------------------------------------
def hatched_regions(segs, k):
    cand = [s for s in segs if T.hatch_len[0] <= s["len"] <= T.hatch_len[1]]
    tree = STRtree([s["g"] for s in cand])
    hatch = []
    for s in cand:
        twins = 0
        for j in tree.query(s["g"].buffer(0.5 * k)):
            t = cand[j]
            if t is s:
                continue
            if abs(math.sin(t["ang"] - s["ang"])) < 0.02 and abs(t["len"] - s["len"]) < 0.15 * s["len"]:
                twins += 1
        if twins >= T.hatch_neighbours:
            hatch.append(s)
    if not hatch:
        return []
    # close the gaps between strokes to recover the hatched area
    blobs = unary_union([s["g"].buffer(0.06 * k) for s in hatch])
    blobs = [b for b in getattr(blobs, "geoms", [blobs])]
    htree = STRtree([s["g"] for s in hatch])
    out = []
    for b in blobs:
        members = [hatch[i] for i in htree.query(b) if b.contains(hatch[i]["g"])]
        if len(members) < T.hatch_min_strokes:
            continue
        fams = collections.Counter(round(math.degrees(s["ang"]) / 5) * 5 % 180 for s in members)
        top = fams.most_common(2)
        cross = (len(top) == 2 and 80 <= abs(top[0][0] - top[1][0]) <= 100
                 and top[1][1] >= 0.3 * top[0][1])
        # thickness: hatch strokes run face to face, so the local spread of their end points
        # across the wall is its thickness
        ends = [pt for s in members for pt in (s["a"], s["b"])]
        etree = STRtree([Point(p) for p in ends])
        samples = []
        for s in members[::max(1, len(members) // 60)]:
            c = s["g"].centroid
            local = [ends[i] for i in etree.query(c.buffer(0.8 * k))]
            if len(local) >= 6:
                mrr = MultiPoint(local).minimum_rotated_rectangle
                if mrr.geom_type == "Polygon":
                    q = list(mrr.exterior.coords)
                    a, bb = math.dist(q[0], q[1]), math.dist(q[1], q[2])
                    if max(a, bb) > 2.5 * min(a, bb):  # a straight stretch, not a corner
                        samples.append(min(a, bb) / k)
        thickness = statistics.median(samples) if samples else None
        region = b.buffer(-0.06 * k).buffer(0)
        region = max(getattr(region, "geoms", [region]), key=lambda g: g.area) if not region.is_empty else b.buffer(0)
        area = region.area / k ** 2
        out.append({"g": region, "strokes": len(members), "cross": cross,
                    "stroke_len": statistics.median(s["len"] for s in members),
                    "family_deg": sorted(a for a, _ in top), "thickness": thickness,
                    "length": area / thickness if thickness else None, "area": area})
    return out


# ---------------------------------------------------------------------------
# 3. Bands: pairs of parallel lines a given width apart
# ---------------------------------------------------------------------------
def band_pieces(segs, width, k, fill_tree, fill_geoms, mode):
    """mode 'empty': nothing filled inside (grade beams); 'wall': a wall runs inside it
    (strip footings under walls, not always centred: perimeter strips can be eccentric)."""
    # element edges: drawn lines (not the border of a fill), solid, not in the axis pen
    long_ = [s for s in segs if s["len"] >= T.band_min_overlap and s["stroke"]
             and not s.get("dashed") and not s.get("axis_pen")]
    tree = STRtree([s["g"] for s in long_])
    pieces = []
    for i, s in enumerate(long_):
        for j in tree.query(s["g"].buffer((width + T.band_tol) * k)):
            if j <= i:
                continue
            t = long_[j]
            if abs(math.sin(t["ang"] - s["ang"])) > 0.005:
                continue
            ang = s["ang"]
            sp = abs(offset(((t["a"][0] + t["b"][0]) / 2, (t["a"][1] + t["b"][1]) / 2), ang) - offset(s["a"], ang)) / k
            if abs(sp - width) > T.band_tol:
                continue
            lo = max(min(along(s["a"], ang), along(s["b"], ang)), min(along(t["a"], ang), along(t["b"], ang)))
            hi = min(max(along(s["a"], ang), along(s["b"], ang)), max(along(t["a"], ang), along(t["b"], ang)))
            if (hi - lo) / k < T.band_min_overlap:
                continue
            t_mid = ((t["a"][0] + t["b"][0]) / 2, (t["a"][1] + t["b"][1]) / 2)
            mid = (offset(s["a"], ang) + offset(t_mid, ang)) / 2
            ux, uy = axis(ang)
            nx, ny = -uy, ux
            half = width * k / 2

            def P(u, v):
                return (u * ux + v * nx, u * uy + v * ny)

            poly = Polygon([P(lo, mid - half), P(hi, mid - half), P(hi, mid + half), P(lo, mid + half)])
            covered = sum(poly.intersection(fill_geoms[q]).area for q in fill_tree.query(poly)) / poly.area
            if mode == "empty" and covered > T.band_empty:
                continue
            if mode == "wall" and covered < T.band_wall_share:
                continue
            pieces.append({"ang": ang, "off": mid, "lo": lo, "hi": hi, "width": width, "poly": poly,
                           "lines": (s["idx"], t["idx"])})
    return pieces


def band_runs(pieces, k):
    """Merge collinear band pieces into continuous runs."""
    groups = collections.defaultdict(list)
    for p in pieces:
        a = round(math.degrees(p["ang"]) * 2) / 2 % 180
        groups[(a, round(p["off"] / (0.05 * k)), p["width"])].append(p)
    runs = []
    for (a, _, w), ps in groups.items():
        ps.sort(key=lambda p: p["lo"])
        cur = None
        for p in ps:
            if cur and p["lo"] <= cur["hi"] + T.band_run_gap * k:
                cur["hi"] = max(cur["hi"], p["hi"])
                cur["pieces"].append(p)
            else:
                cur = {"ang": p["ang"], "off": p["off"], "lo": p["lo"], "hi": p["hi"], "width": w, "pieces": [p]}
                runs.append(cur)
    for r in runs:
        ux, uy = axis(r["ang"])
        nx, ny = -uy, ux
        P = lambda u, v: (u * ux + v * nx, u * uy + v * ny)
        half = r["width"] * k / 2
        r["poly"] = Polygon([P(r["lo"], r["off"] - half), P(r["hi"], r["off"] - half),
                             P(r["hi"], r["off"] + half), P(r["lo"], r["off"] + half)])
        r["axis"] = LineString([P(r["lo"], r["off"]), P(r["hi"], r["off"])])
        r["length"] = (r["hi"] - r["lo"]) / k
    # a run shorter than a typical band piece and fully inside another run is a duplicate
    runs.sort(key=lambda r: -r["length"])
    kept = []
    for r in runs:
        if any(q["poly"].buffer(0.02 * k).contains(r["poly"]) for q in kept):
            continue
        kept.append(r)
    return kept


def pick_band_width(labels, candidates, segs, k, fill_tree, fill_geoms, mode):
    """Which number written in the labels is the band width? The one that has a band of that
    width next to most of the labels."""
    best = None
    for w in candidates:
        pieces = band_pieces(segs, w, k, fill_tree, fill_geoms, mode)
        if not pieces:
            continue
        tree = STRtree([p["poly"] for p in pieces])
        n = sum(1 for l in labels if len(tree.query(l["c"].buffer(T.label_radius * k))))
        if best is None or n > best[1]:
            best = (w, n, pieces)
    return best


def junctions(runs, k):
    """Points where the axes of two non-parallel runs meet (T, L or X junctions)."""
    pts = []
    for i, r in enumerate(runs):
        for q in runs[i + 1:]:
            if abs(math.sin(r["ang"] - q["ang"])) < math.sin(math.radians(20)):
                continue
            ext = (max(r["width"], q["width"]) / 2 + 0.3) * k
            a = r["axis"]
            b = q["axis"]
            # extend both axes a little so a run stopping at the other's edge still meets it
            p = _line_intersection(a, b)
            if p and a.distance(p) <= ext and b.distance(p) <= ext:
                pts.append(p)
    merged = []
    for p in pts:
        if all(p.distance(m) > T.junction_radius * k for m in merged):
            merged.append(p)
    return merged


def _line_intersection(a, b):
    (x1, y1), (x2, y2) = a.coords[0], a.coords[-1]
    (x3, y3), (x4, y4) = b.coords[0], b.coords[-1]
    d = (x1 - x2) * (y3 - y4) - (y1 - y2) * (x3 - x4)
    if abs(d) < 1e-9:
        return None
    t = ((x1 - x3) * (y3 - y4) - (y1 - y3) * (x3 - x4)) / d
    return Point(x1 + t * (x2 - x1), y1 + t * (y2 - y1))


# ---------------------------------------------------------------------------
# 4. Grid axes: long chains of collinear dashes, with a bubble label at an end
# ---------------------------------------------------------------------------
def bubbles(paths, texts):
    """Grid bubbles: small circles (curves only) with a short name inside."""
    circles = []
    for p in paths:
        if p["type"] != "s" or not p["items"] or any(kind != "c" for kind, _ in p["items"]):
            continue
        x0, y0, x1, y1 = p["rect"]
        w, h = x1 - x0, y1 - y0
        if w > 0 and 0.85 < h / w < 1.15:
            circles.append((Point((x0 + x1) / 2, (y0 + y1) / 2), w / 2))
    names = [t for t in texts if re.fullmatch(r"[A-Z]{1,2}\d?|\d{1,2}[A-Z]?", t["text"])]
    out = []
    for c, r in circles:
        inside = [t for t in names if ds.centre(t["rect"]).distance(c) < 0.6 * r]
        if inside:
            out.append({"c": c, "r": r, "tag": inside[0]["text"]})
    return out


def grid_lines(paths, texts, k, tboxes):
    dashes = straight_segments(paths, k, tboxes, 0.02, 2.0)
    groups = collections.defaultdict(list)
    for s_ in dashes:
        a = round(math.degrees(s_["ang"]) * 5) / 5 % 180
        mid = ((s_["a"][0] + s_["b"][0]) / 2, (s_["a"][1] + s_["b"][1]) / 2)
        groups[(a, round(offset(mid, s_["ang"]) / (0.03 * k)))].append(s_)
    bub = bubbles(paths, texts)
    lines = []
    for (a, _), ss in groups.items():
        if len(ss) < T.grid_min_pieces:
            continue
        ang = ss[0]["ang"]
        iv = sorted((min(along(x["a"], ang), along(x["b"], ang)), max(along(x["a"], ang), along(x["b"], ang))) for x in ss)
        lo, hi = iv[0][0], max(v[1] for v in iv)
        extent = (hi - lo) / k
        cover = sum(v[1] - v[0] for v in iv) / (hi - lo) if hi > lo else 1
        if extent < T.grid_min_extent or cover > 0.97:   # a continuous line is not a dashed axis
            continue
        off = statistics.median(offset(((x["a"][0] + x["b"][0]) / 2, (x["a"][1] + x["b"][1]) / 2), ang) for x in ss)
        ux, uy = axis(ang)
        P = lambda u: Point(u * ux - off * uy, u * uy + off * ux)
        ends = (P(lo), P(hi))
        tag = None
        for b in bub:
            # the bubble sits on the axis' extension, just beyond one of its ends
            if abs(offset((b["c"].x, b["c"].y), ang) - off) < b["r"] * 0.5 and \
                    min(b["c"].distance(e) for e in ends) < b["r"] + 3.0 * k:
                tag = b["tag"]
                break
        if tag:
            lines.append({"ang": ang, "off": off, "line": LineString([ends[0], ends[1]]), "tag": tag, "extent": extent})
    uniq = []
    for l in sorted(lines, key=lambda l: -l["extent"]):
        if not any(abs(math.sin(l["ang"] - u["ang"])) < 0.005 and abs(l["off"] - u["off"]) < 0.05 * k for u in uniq):
            uniq.append(l)
    return uniq


# ---------------------------------------------------------------------------
# 5. Closed outlines (footings, pads)
# ---------------------------------------------------------------------------
def closed_outlines(paths, k, tboxes):
    out = []
    for i, p in enumerate(paths):
        if p["type"] != "s":
            continue
        for g in ds.subpolygons(p):
            if g.area / k ** 2 < 0.1 or in_text(g, tboxes):
                continue
            o = ds.oriented(g, k)
            if o["rectness"] > 0.97:
                out.append({"g": g, "o": o, "idx": i, "width_pt": p["width"]})
    kept = []
    for r in sorted(out, key=lambda r: -r["g"].area):
        if not any(abs(r["g"].area - q["g"].area) / q["g"].area < 0.02 and r["g"].intersection(q["g"]).area / r["g"].area > 0.98
                   for q in kept):
            kept.append(r)
    return kept


# ---------------------------------------------------------------------------
# 6. Text-driven elements: rafts, sump pits
# ---------------------------------------------------------------------------
def raft_stamps(texts, segs, k):
    """'Radier' stamp + the number under it (thickness, cm). The raft's extent comes from the
    diagonals drawn from its corners towards the stamp: each stops at the stamp circle, so the
    far ends of those rays are the raft's corners."""
    nums = [t for t in texts if re.fullmatch(r"\d{2,3}", t["text"])]
    long_ = [s_ for s_ in segs if s_["len"] >= 1.5]
    tree = STRtree([s_["g"] for s_ in long_])
    out = []
    for t in texts:
        if not re.match(r"(?i)^radier", t["text"]):
            continue
        c = ds.centre(t["rect"])
        m = re.search(r"(\d{2,3})\s*CM", t["text"])
        thick = int(m.group(1)) if m else None
        if thick is None and nums:
            n = min(nums, key=lambda n: ds.centre(n["rect"]).distance(c))
            if ds.centre(n["rect"]).distance(c) < 1.5 * k:     # right under the word
                thick = int(n["text"])
        corners, dirs = [], set()
        for i in tree.query(c.buffer(1.5 * k)):
            s_ = long_[i]
            ux, uy = axis(s_["ang"])
            perp = abs(-(c.x - s_["a"][0]) * uy + (c.y - s_["a"][1]) * ux) / k
            da, db = math.dist(s_["a"], (c.x, c.y)) / k, math.dist(s_["b"], (c.x, c.y)) / k
            if perp < 0.4 and min(da, db) < 1.2 and max(da, db) > 2.0:
                far = s_["b"] if db > da else s_["a"]
                corners.append(far)
                dirs.add(round(math.degrees(math.atan2(far[1] - c.y, far[0] - c.x)) / 20))
        quad = None
        if len(dirs) >= 3:
            hull = MultiPoint(corners).convex_hull
            if hull.geom_type == "Polygon":
                quad = hull
        out.append({"c": c, "text": t["text"], "thickness_cm": thick, "g": quad, "rays": len(corners)})
    return out


def sump_pits(texts):
    out = []
    for t in texts:
        m = re.search(r"(?i)fosse\s+de\s+relevage.*?(\d+[.,]\d+)\s*x\s*(\d+[.,]\d+)\s*x\s*(\d+[.,]\d+)", t["text"])
        if m:
            dims = [float(v.replace(",", ".")) for v in m.groups()]
            out.append({"c": ds.centre(t["rect"]), "dims": dims, "text": t["text"]})
    return out


# ---------------------------------------------------------------------------
# 7. Labels and buildings
# ---------------------------------------------------------------------------
def all_labels(texts, cfg):
    labels = ds.parse_labels(texts, cfg)          # column, wall, retaining_wall, footing
    for kind, r in cfg["bands"].items():
        rx = re.compile(r["label"])
        for t in texts:
            m = rx.search(t["text"])
            if m:
                g = {n: v or "" for n, v in m.groupdict().items()}
                labels.append({"kind": kind, "raw": t["text"], "c": ds.centre(t["rect"]), "key": r["key"].format(**g),
                               "groups": g, "building": None, "dims": None,
                               "numbers": [int(g["a"]) / 100, int(g["b"]) / 100]})
    return labels


class Buildings:
    """Building of a point: majority of the nearest labels that carry a building number."""

    def __init__(self, labels, k):
        self.ls = [l for l in labels if l.get("building")]
        self.k = k

    def of(self, pt):
        near = sorted(self.ls, key=lambda l: l["c"].distance(pt))[:T.building_k]
        near = [l for l in near if l["c"].distance(pt) < 25 * self.k]
        if not near:
            return None, 0.0
        b, n = collections.Counter(l["building"] for l in near).most_common(1)[0]
        return b, n / len(near)


# ---------------------------------------------------------------------------
# 8. Elements
# ---------------------------------------------------------------------------
def element(family, g, k, rule, **extra):
    o = ds.oriented(g, k) if g.geom_type == "Polygon" else None
    e = {"family": family, "subclass": family, "g": g, "geom": g, "o": o, "rule": rule, "label": None,
         "label_source": "none", "marks": {}, "ref": None, "flags": []}
    e.update(extra)
    return e


def wall_piece(obj, walls, k):
    """A short filled piece touching a wall of the same thickness (in line or at a corner) is a
    piece of that wall, not a column: on the verified plan no column touches a filled wall."""
    o = obj["o"]
    for w in walls:
        if w is obj or obj["g"].distance(w["g"]) > T.touch * k:
            continue
        if abs(w["o"]["width_m"] - o["width_m"]) <= 0.015:
            return True
    return False


def match_ref(g, refs):
    best = (0.0, None)
    for r in refs:
        if r["g"].intersects(g):
            iou = r["g"].intersection(g).area / r["g"].union(g).area
            if iou > best[0]:
                best = (iou, r)
    return best[1] if best[0] > 0.9 else None


def load_reference(path, k_px):
    """Verified detections from detect_structures.py (these used colour; here they only grade)."""
    if not Path(path).exists():
        return []
    r = json.loads(Path(path).read_text())
    s = 72 / r["image"]["dpi"]
    out = []
    for d in r["detections"]:
        pts = d.get("polygon_px") or d["obb_px"]
        out.append({"subclass": d["subclass"], "class": d["class"], "type_key": d["type_key"],
                    "g": Polygon([(x * s, y * s) for x, y in pts]).buffer(0)})
    return out


def path_style(paths, idx):
    p = paths[idx]
    col = p["fill"] or p["color"]
    return (list(col) if col else None), p["layer"]


def build(pdf, cfg, ref_path):
    doc, page, paths, texts = ds.load_page(pdf)
    glyph = ds.glyph_size(texts, cfg)
    calib = ds.calibrate(paths, texts, glyph, cfg)
    k = calib["pt_per_m"]
    ctx = SimpleNamespace(k=k, cfg=cfg, glyph=glyph)
    tb = text_index(texts)
    segs = straight_segments(paths, k, tb)
    dashed, axis_pen = mark_dashed(segs, k)
    labels = all_labels(texts, cfg)
    kinds = collections.defaultdict(list)
    for l in labels:
        kinds[l["kind"]].append(l)
    blds = Buildings(labels, k)
    refs = load_reference(ref_path, k)
    log = {"calibration": calib, "dashed_pieces": dashed, "axis_pen_pt": axis_pen}

    # --- filled objects -> columns and walls ------------------------------------------
    rects, dups, others = filled_objects(paths, k)
    log["filled"] = {"rectangular_objects": len(rects), "drawn_twice": dups, "other_filled_objects": len(others)}
    fill_geoms = [r["g"] for r in rects] + [g for g in others if g.area > 0.02 * k * k]
    fill_tree = STRtree(fill_geoms)
    for r in rects:
        r["ec2"] = ec2_class(r["o"])
    # thin bent fills (walls drawn around corners as one object) -> straight stretches
    bent = 0
    for g in others:
        area = g.area / k ** 2
        if area < 0.5 or g.geom_type != "Polygon":
            continue
        if 2 * area / (g.length / k) > 0.4:      # not thin: mean thickness 2A/P above 40 cm
            continue
        for piece in straight_stretches(g, k):
            if any(piece.intersection(r["g"]).area > 0.5 * piece.area for r in rects if not r.get("bent")):
                continue      # already found as its own rectangle
            oo = ds.oriented(piece, k)
            rects.append({"g": piece, "o": oo, "paths": [], "ec2": ec2_class(oo), "bent": True, "src": g})
            bent += 1
    log["filled"]["stretches_from_bent_walls"] = bent
    wallish = [r for r in rects if r["ec2"] == "wall"]
    elems = []
    for r in rects:
        fam = r["ec2"]
        rule = "solid fill, rectangle, L <= 4 W (EN 1992-1-1 5.3.1(7))" if fam == "column" else \
            "solid fill, rectangle, L > 4 W (EN 1992-1-1 5.3.1(7))"
        if fam == "column" and (r.get("bent") or wall_piece(r, wallish, k)):
            fam, rule = "wall", "solid fill, short piece touching a wall of the same thickness"
        if r.get("bent"):
            rule = "straight stretch of a thin filled outline with bends (pair of parallel faces)"
        e = element(fam, r["g"], k, rule, pieces=len(r["paths"]), ec2=r["ec2"],
                    style=path_style(paths, r["paths"][0]) if r["paths"] else (None, None))
        elems.append(e)

    # --- cross-hatched walls ---------------------------------------------------------
    hatch = hatched_regions(segs, k)
    log["hatched_regions"] = {"cross": sum(h["cross"] for h in hatch), "other_patterns": sum(not h["cross"] for h in hatch)}
    for h in hatch:
        if h["cross"] and h["length"] and h["length"] >= 0.3:
            elems.append(element("core_wall", h["g"], k, "cross-hatch of short strokes in two perpendicular families",
                                 hatch=h, style=(None, None)))
    hatched = [e for e in elems if e["family"] == "core_wall"]

    # --- bands: grade beams and strip footings ------------------------------------------
    walls_for_bands = fill_geoms + [e["g"] for e in hatched]
    wall_tree = STRtree(walls_for_bands)
    runs_by_kind = {}
    for kind, r in cfg["bands"].items():
        groups = collections.defaultdict(list)
        for l in kinds[kind]:
            groups[l["key"]].append(l)
        runs_by_kind[kind] = []
        for key, ls in groups.items():
            tree, geoms = (fill_tree, fill_geoms) if r["mode"] == "empty" else (wall_tree, walls_for_bands)
            best = pick_band_width(ls, ls[0]["numbers"], segs, k, tree, geoms, r["mode"])
            if not best:
                continue
            width, support, pieces = best
            log.setdefault("bands", {})[key] = {"labels": len(ls), "width_m": width, "labels_with_band": support}
            for run in band_runs(pieces, k):
                run["key"], run["kind"] = key, kind
                runs_by_kind[kind].append(run)
    # a run claimed by two keys of the same kind keeps the better-supported one
    for kind, runs in runs_by_kind.items():
        kept = []
        for run in sorted(runs, key=lambda r: -r["length"]):
            if any(q["poly"].buffer(0.02 * k).contains(run["poly"]) for q in kept):
                continue
            kept.append(run)
        runs_by_kind[kind] = kept

    # --- closed outlines -> footings, pads -----------------------------------------------
    outlines = closed_outlines(paths, k, tb)
    rafts = raft_stamps(texts, segs, k)
    pits = sump_pits(texts)
    tiny = [ds.centre(p["rect"]) for p in paths if max(p["rect"][2] - p["rect"][0], p["rect"][3] - p["rect"][1]) < 0.1 * k]
    tiny_tree = STRtree(tiny)
    solid = [e for e in elems if e["family"] in ("column", "wall")]
    solid_tree = STRtree([e["g"] for e in solid])
    flabels = kinds["footing"]
    for o in outlines:
        g, oo = o["g"], o["o"]
        area = g.area / k ** 2
        side_ok = T.pad_side[0] <= oo["width_m"] and oo["length_m"] <= T.pad_side[1] and oo["length_m"] / oo["width_m"] < 1.2
        dots = sum(1 for i in tiny_tree.query(g) if g.contains(tiny[i]))
        if side_ok and dots >= T.pad_min_dots:
            elems.append(element("mass_concrete_pad", g, k, "small closed square filled with stipple dots, no label",
                                 dots=dots, style=path_style(paths, o["idx"])))
            continue
        if not (T.footing_min_area <= area <= 60 and oo["length_m"] / oo["width_m"] <= T.footing_max_aspect):
            continue
        if any(g.contains(r["c"]) for r in rafts):
            continue  # a raft outline, handled with the raft
        sup = [solid[i] for i in solid_tree.query(g) if solid[i]["g"].intersection(g).area > 0.5 * solid[i]["g"].area]
        lab = [l for l in flabels if g.contains(l["c"]) or g.exterior.distance(l["c"]) < 1.0 * k]
        if sup or lab:
            elems.append(element("footing", g, k, "closed rectangular outline around a column/wall or next to an S/SC label",
                                 supports=len(sup), style=path_style(paths, o["idx"])))
    footings = [e for e in elems if e["family"] == "footing"]
    # nested duplicates: an outline inside another footing with the same supports is its inner line
    keep = []
    for f in sorted(footings, key=lambda f: -f["g"].area):
        if any(q["g"].contains(f["g"]) and q["supports"] == f["supports"] and f["supports"] > 0 for q in keep):
            f["drop"] = True
        else:
            keep.append(f)
    elems = [e for e in elems if not e.get("drop")]
    footings = [e for e in elems if e["family"] == "footing"]
    raft_extents = [r["g"] for r in rafts if r["g"] is not None]

    # --- labels -> shapes ------------------------------------------------------------------
    columns = [e for e in elems if e["family"] == "column"]
    walls = [e for e in elems if e["family"] == "wall"]
    for e in elems:
        e["label"] = None
    ds.assign(columns, kinds["column"], T.label_radius * k, ds.near_geom)
    # size-aware: pass 1 learns each V type's size (per building), pass 2 lets a label only onto a
    # wall of its type's size, so a V label next to a long wall run goes to its own short wall
    left_v = ds.assign_by_catalogue(walls, kinds["wall"], SimpleNamespace(k=k, cfg=cfg), ds.near_geom)
    left_vp = []    # VP labels name the wall run by its thickness: checked below, not assigned one-to-one
    left_f = ds.assign_by_catalogue(footings, flabels, SimpleNamespace(k=k, cfg=cfg), ds.near_outline)
    for group in (columns, walls, footings):
        for e in group:
            if e["label"]:
                e["label_source"] = "printed"
    ctx_i = SimpleNamespace(k=k, cfg=cfg)
    ds.infer_missing_labels(columns, labels, cfg["labels"]["column"], ctx_i)
    ds.infer_missing_labels([w for w in walls if not w["label"] or w["label"]["kind"] == "wall"], labels, cfg["labels"]["wall"], ctx_i)
    ds.infer_missing_labels(footings, labels, cfg["labels"]["footing"], ctx_i)
    for e in columns + walls + footings:
        if e["label"] and e["label"].get("inferred"):
            e["label_source"] = "inferred from building + size"
    repeats, unmatched = [], []
    for l in left_v + left_vp + left_f:
        pool = walls if l["kind"] in ("wall", "retaining_wall") else footings
        near = [e for e in pool if e["label"] and e["label"]["key"] == l["key"]
                and e["g"].distance(l["c"]) < T.label_radius * k]
        (repeats if near else unmatched).append(l["raw"])
    log["labels_repeating"] = repeats
    log["labels_without_shape"] = unmatched

    # --- subclasses -----------------------------------------------------------------------
    sf_polys = [r["poly"] for r in runs_by_kind.get("strip_footing", [])]
    series = cfg["labels"]["footing"]["series_subclass"]
    vp_rule = cfg["labels"]["retaining_wall"]
    vp_numbers = sorted({int(l["groups"]["type"]) for l in kinds["retaining_wall"]})
    sf_union = unary_union(sf_polys) if sf_polys else None
    for w in walls:
        # a perimeter wall usually spans several strip-footing runs (they break at footings)
        in_sf = sf_union is not None and sf_union.intersection(w["g"]).area > 0.5 * w["g"].area
        t_cm = round(w["o"]["width_m"] * 100)
        by_vp = any(w["g"].distance(l["c"]) < T.label_radius * k and int(l["groups"]["type"]) == t_cm
                    for l in kinds["retaining_wall"])
        if (in_sf or by_vp) and not (w["label"] and w["label"]["kind"] == "wall"):
            w["subclass"] = "retaining_wall"
            t = round(w["o"]["width_m"] * 100)
            if t in vp_numbers:     # VP20 = 20 cm thick: the grammar names the type by thickness
                w["label"] = {"raw": None, "key": vp_rule["key"].format(type=t), "kind": "retaining_wall",
                              "groups": {"type": str(t)}}
                w["label_source"] = "thickness matches a VP type"
        else:
            w["subclass"] = "shear_wall"
    # every printed VP label should sit by a retaining wall of that thickness
    ret = [w for w in walls if w["subclass"] == "retaining_wall"]
    vp_check = collections.Counter()
    for l in kinds["retaining_wall"]:
        near = [w for w in ret if w["g"].distance(l["c"]) < T.label_radius * k]
        if not near:
            vp_check["no retaining wall within 1.5 m"] += 1
        elif any(round(w["o"]["width_m"] * 100) == int(l["groups"]["type"]) for w in near):
            vp_check["thickness agrees"] += 1
        else:
            vp_check["thickness differs"] += 1
    log["vp_labels"] = dict(vp_check)
    for f in footings:
        if not f["label"] and any(f["g"].intersection(r).area > 0.5 * f["g"].area for r in raft_extents):
            f["drop"] = True        # an unlabelled outline mostly inside a raft is part of the raft
            continue
        if f["label"]:
            f["subclass"] = series.get(f["label"]["groups"].get("series"), "isolated_footing")
        else:
            f["subclass"] = "combined_footing" if f["supports"] >= 2 else "isolated_footing"
            f["ref_note"] = "closed outline around a wall/column with no footing label"
        if f["label"] and f["label"].get("dims"):
            notes = [ds.centre(t["rect"]) for t in texts if re.search(cfg["crane"]["note_pattern"], t["text"], re.I)]
            if any(f["g"].distance(n) < cfg["crane"]["group_radius_m"] * k for n in notes):
                f["subclass"] = "crane_footing"
    elems = [e for e in elems if not e.get("drop")]
    footings = [e for e in elems if e["family"] == "footing"]
    cranes = [f for f in footings if f["subclass"] == "crane_footing"]
    if len(cranes) >= cfg["crane"]["min_members"]:
        hull = unary_union([f["g"] for f in cranes]).convex_hull
        elems.append(element("crane_foundation", hull, k, "footings of stated size grouped next to a crane note",
                             members=len(cranes), style=(None, None)))

    # --- rafts, pits, bands as elements --------------------------------------------------
    for r in rafts:
        near_pit = any(r["c"].distance(p["c"]) < 4 * k for p in pits)
        if near_pit:
            continue
        g = r["g"] or r["c"].buffer(0.5 * k)
        e = element("raft", g, k, "'Radier' stamp; extent from the corner diagonals that stop at the stamp",
                    thickness=r["thickness_cm"], rays=r["rays"], extent_found=r["g"] is not None, style=(None, None))
        if r["g"] is None:
            e["o"] = None           # stamp found, extent not: no size reported
            e["rule"] = "'Radier' stamp; extent not found (no corner diagonals)"
        e["label"] = {"raw": r["text"], "key": f"Radier {r['thickness_cm']}", "kind": "raft"}
        e["label_source"] = "printed"
        if r["thickness_cm"] and r["thickness_cm"] >= 150:
            e["subclass"] = "crane_raft (probable)"
        elems.append(e)
    for p in pits:
        e = element("sump_pit", p["c"].buffer(0.5 * k), k, "'FOSSE DE RELEVAGE LxWxH' text (size read from the text)",
                    dims=p["dims"], style=(None, None))
        e["label"] = {"raw": p["text"], "key": "Fosse {:.2f}x{:.2f}x{:.2f}".format(*p["dims"]), "kind": "pit"}
        e["label_source"] = "printed"
        elems.append(e)
    # grade-beam spans: a run split wherever it passes a footing, wall, column or pad
    supports = [e["g"] for e in elems if e["family"] in ("footing", "wall", "column", "mass_concrete_pad", "core_wall")]
    sup_tree = STRtree(supports)
    for kind, runs in runs_by_kind.items():
        for run in runs:
            if kind == "grade_beam":
                cut = unary_union([supports[i] for i in sup_tree.query(run["poly"])])
                rest = run["axis"].difference(cut) if not cut.is_empty else run["axis"]
                parts = [p for p in getattr(rest, "geoms", [rest]) if p.length / k >= 0.3]
                for p in parts:
                    poly = p.buffer(run["width"] * k / 2, cap_style=2)
                    e = element("grade_beam", poly, k, "two parallel lines at the width written in the LG label, empty inside",
                                span_m=p.length / k, band_width=run["width"], style=(None, None))
                    e["label"] = {"raw": None, "key": run["key"], "kind": kind}
                    e["label_source"] = "band width matches label"
                    e["lines"] = run["pieces"][0]["lines"]
                    elems.append(e)
            else:
                e = element("strip_footing", run["poly"], k, "two parallel lines at the width written in the SF label, a wall inside",
                            span_m=run["length"], band_width=run["width"], style=(None, None))
                e["label"] = {"raw": None, "key": run["key"], "kind": kind}
                e["label_source"] = "band width matches label"
                e["lines"] = run["pieces"][0]["lines"]
                elems.append(e)
    lg_runs = runs_by_kind.get("grade_beam", [])
    juncs = junctions(lg_runs, k)
    grid = grid_lines(paths, texts, k, tb)

    # --- colour-free marks ------------------------------------------------------------------
    fam_geoms = lambda fams: [e for e in elems if e["family"] in fams]
    foot_list = fam_geoms(("footing",))
    raft_list = [e for e in elems if e["family"] == "raft" and e["extent_found"]]
    lg_polys = [r["poly"] for r in lg_runs]
    lg_tree = STRtree(lg_polys) if lg_polys else None
    j_tree = STRtree(juncs) if juncs else None
    hatched_g = [e["g"] for e in hatched]
    solid_all = fam_geoms(("column", "wall"))
    for e in elems:
        g, m = e["g"], e["marks"]
        c = g.centroid
        cont = [f for f in foot_list if f is not e and f["g"].buffer(0.02 * k).intersection(g).area > 0.5 * g.area]
        m["inside_footing"] = cont[0]["label"]["key"] if cont and cont[0]["label"] else ("unlabelled" if cont else "")
        m["on_raft"] = "yes" if any(r["g"].buffer(0.1 * k).contains(c) for r in raft_list if r is not e) else "no"
        m["touches_hatched_wall"] = "yes" if any(h.distance(g) < 0.05 * k for h in hatched_g if h is not g) else "no"
        m["touches_solid_element"] = "yes" if any(s["g"].distance(g) < T.touch * k for s in solid_all if s is not e) else "no"
        m["in_strip_band"] = "yes" if sf_union is not None and sf_union.intersection(g).area > 0.5 * g.area else "no"
        m["on_grade_beam"] = "yes" if lg_tree is not None and any(lg_polys[i].intersects(g) for i in lg_tree.query(g)) \
            and e["family"] != "grade_beam" else "no"
        m["at_lg_junction"] = "yes" if j_tree is not None and len(j_tree.query(g.buffer(0.3 * k))) else "no"
        m["grid_distance_m"] = round(min((gl["line"].distance(c) for gl in grid), default=float("nan")) / k, 2)
        if e["family"] == "footing":
            inside = [s for s in solid_all if g.buffer(0.02 * k).intersection(s["g"]).area > 0.5 * s["g"].area]
            m["supports"] = "+".join(sorted(collections.Counter(s["family"] for s in inside).elements())) or "none"
            if inside:
                sc = unary_union([s["g"] for s in inside]).centroid
                m["support_offset_cm"] = round(sc.distance(c) / k * 100)
        e["building"], e["building_conf"] = (e["label"]["groups"]["bld"], 1.0) \
            if e["label"] and e["label"].get("groups", {}).get("bld") else blds.of(c)
        e["building_source"] = "own label" if e["building_conf"] == 1.0 and e["label"] and \
            e["label"].get("groups", {}).get("bld") else f"nearest labels ({e['building_conf']:.0%} agree)"

    # outside the drawing: no labelled element within reach (title block tables, legends)
    outside = [e for e in elems if e["building"] is None]
    log["outside_plan"] = dict(collections.Counter(e["family"] for e in outside))
    elems = [e for e in elems if e["building"] is not None]
    for e in elems:
        if e["family"] == "wall" and e["subclass"] == "shear_wall" and not e["label"] and e["marks"]["on_raft"] == "yes":
            e["subclass"] = "core_wall_solid"   # lift/stair cores stand on rafts and carry no V mark here

    # --- what sits at each grade-beam junction ----------------------------------------------------
    # an element within 0.6 m (a beam width plus margin) of the crossing point; otherwise the
    # contents of the footing the crossing falls in (walls often stand eccentric on the footing)
    jrows = []
    order = ("column", "wall", "core_wall", "mass_concrete_pad")
    for p in juncs:
        near = [e for e in elems if e["family"] in order and e["g"].distance(p) < 0.6 * k]
        fams = {e["family"] for e in near}
        what = next((f for f in order if f in fams), None)
        foot = next((f for f in elems if f["family"] == "footing" and f["g"].contains(p)), None)
        if what is None and foot is not None:
            sup = foot["marks"].get("supports", "none")
            what = "footing carrying " + ("wall" if "wall" in sup else "column" if "column" in sup else "nothing")
        if what is None and any(e["family"] == "raft" and e["extent_found"] and e["g"].contains(p) for e in elems):
            what = "raft"
        b, conf = blds.of(p)
        if b is None:
            continue    # outside the drawing area
        wall_key = next((e["label"]["key"] for e in near if e["family"] == "wall" and e["label"]), "")
        cands = [e for e in elems if e["family"] in order]
        nearest = min(cands, key=lambda e: e["g"].distance(p)) if cands else None
        jrows.append({"x": p.x, "y": p.y, "building": b, "what": what or "nothing", "wall_key": wall_key,
                      "all": "+".join(sorted(fams | ({"footing"} if foot else set()))),
                      "nearest": nearest["family"] if nearest else "",
                      "nearest_m": round(nearest["g"].distance(p) / k, 2) if nearest else None})

    # --- bands graded by the layer of their lines (validation only) ----------------------------
    for kind, runs in runs_by_kind.items():
        idx = {i for r in runs for pc in r["pieces"] for i in pc["lines"]}
        lay = collections.Counter(paths[i]["layer"] for i in idx)
        log.setdefault("band_layers", {})[kind] = {"lines": len(idx), "top_layers": lay.most_common(4),
                                                   "total_run_m": round(sum(r["length"] for r in runs), 1),
                                                   "runs": len(runs)}

    # --- grade against the verified detections -----------------------------------------------
    ref_map = {"column": ("column",), "wall": ("shear_wall", "retaining_wall"),
               "footing": ("isolated_footing", "combined_footing", "crane_footing"), "mass_concrete_pad": ("mass_concrete_pad",)}
    grading = []
    for fam, subs in ref_map.items():
        rs = [r for r in refs if r["subclass"] in subs]
        mine = [e for e in elems if e["family"] == fam]
        for e in mine:
            e["ref"] = match_ref(e["g"], rs)
        other = [e for e in elems if e["family"] != fam]
        found = sum(1 for r in rs if any(e["ref"] is r for e in mine))
        wrong_family = sum(1 for r in rs if any(match_ref(e["g"], [r]) for e in other if e["g"].intersects(r["g"])))
        grading.append({"family": fam, "verified": len(rs), "found_colour_free": found,
                        "found_as_other_family": wrong_family, "colour_free_total": len(mine),
                        "not_in_verified_set": sum(1 for e in mine if e["ref"] is None)})
    nets = [r["g"] for r in refs if r["subclass"] == "wall_network"]
    for e in elems:
        if e["family"] == "wall" and e["ref"] is None and any(n.buffer(0.05 * k).contains(e["g"]) for n in nets):
            e["ref_note"] = "piece of a wall network the colour pipeline left unsplit"
    return {"elems": elems, "junctions": jrows, "grading": grading, "log": log, "k": k, "grid": grid,
            "page": page, "paths": paths, "calib": calib}


# ---------------------------------------------------------------------------
# 9. Workbook
# ---------------------------------------------------------------------------
DRAWING = {"column": "solid fill", "wall": "solid fill", "core_wall": "cross-hatch", "footing": "closed outline",
           "mass_concrete_pad": "closed outline + stipple", "grade_beam": "parallel line pair",
           "strip_footing": "parallel line pair", "raft": "stamp + corner diagonals", "sump_pit": "text",
           "crane_foundation": "group of footings"}

ELEMENT_COLS = [
    ("ID", 6), ("Building", 9), ("Building source", 22), ("Family", 17), ("Subclass", 19), ("Type key", 14),
    ("Label text", 20), ("Label source", 22), ("Length (cm)", 10), ("Width (cm)", 10), ("L/W", 7), ("Area (m²)", 9),
    ("Span / run (m)", 10), ("Angle (°)", 8), ("Drawing type", 18), ("Colour-free rule", 48), ("EC2 test", 9),
    ("Inside footing", 13), ("Supports inside", 16), ("Support offset (cm)", 10), ("On raft", 8),
    ("Touches hatched wall", 10), ("Touches solid element", 10), ("In strip band", 9), ("On grade beam", 9),
    ("At LG junction", 9), ("Grid distance (m)", 9), ("Hatch stroke (m)", 9), ("Stipple dots", 8),
    ("Raft thickness (cm)", 9), ("bbox x0 (px)", 9), ("bbox y0 (px)", 9), ("bbox x1 (px)", 9), ("bbox y1 (px)", 9),
    ("Matches verified detection", 11), ("Verified subclass", 16), ("Note", 36),
    ("(validation only) colour", 14), ("(validation only) layer", 16),
]


def col_letter(name):
    from openpyxl.utils import get_column_letter
    return get_column_letter([c for c, _ in ELEMENT_COLS].index(name) + 1)


def element_rows(res, dpi):
    k, s = res["k"], dpi / 72
    rows = []
    fam_order = ["column", "wall", "core_wall", "footing", "mass_concrete_pad", "grade_beam", "strip_footing",
                 "raft", "sump_pit", "crane_foundation"]
    elems = sorted(res["elems"], key=lambda e: (str(e["building"]), fam_order.index(e["family"]), e["subclass"],
                                                 e["label"]["key"] if e["label"] else "~"))
    for n, e in enumerate(elems, 1):
        o, m, g = e["o"], e["marks"], e["g"]
        x0, y0, x1, y1 = g.bounds
        L = round(o["length_m"] * 100) if o and e["family"] not in ("sump_pit",) else None
        W = round(o["width_m"] * 100) if o and e["family"] not in ("sump_pit",) else None
        area = round(g.area / k ** 2, 3) if (o or e["family"] == "core_wall") else None
        if e["family"] == "sump_pit":
            L, W = round(e["dims"][0] * 100), round(e["dims"][1] * 100)
            area = round(e["dims"][0] * e["dims"][1], 3)
        if e["family"] == "core_wall":
            h = e["hatch"]
            L, W = (round(h["length"] * 100) if h["length"] else None), (round(h["thickness"] * 100) if h["thickness"] else None)
        if e["family"] in ("grade_beam", "strip_footing"):
            L, W = round(e["span_m"] * 100), round(e["band_width"] * 100)
        lab = e["label"]
        ref = e["ref"]
        family, subclass = e["family"], e["subclass"]
        if family == "core_wall":
            family, subclass = "wall", "core_wall_hatched"
        rows.append({
            "ID": n, "Building": e["building"] or "unknown", "Building source": e["building_source"],
            "Family": family, "Subclass": subclass, "Type key": lab["key"] if lab else "(none)",
            "Label text": (lab.get("raw") or "") if lab else "", "Label source": e["label_source"],
            "Length (cm)": L, "Width (cm)": W, "L/W": None, "Area (m²)": area,
            "Span / run (m)": round(e["span_m"], 2) if "span_m" in e else (round(e["hatch"]["length"], 2) if e["family"] == "core_wall" and e["hatch"]["length"] else None),
            "Angle (°)": round(o["angle_deg"], 1) if o else None, "Drawing type": DRAWING[e["family"]],
            "Colour-free rule": e["rule"], "EC2 test": None,
            "Inside footing": m.get("inside_footing", ""), "Supports inside": m.get("supports", ""),
            "Support offset (cm)": m.get("support_offset_cm"), "On raft": m["on_raft"],
            "Touches hatched wall": m["touches_hatched_wall"], "Touches solid element": m["touches_solid_element"],
            "In strip band": m["in_strip_band"], "On grade beam": m["on_grade_beam"], "At LG junction": m["at_lg_junction"],
            "Grid distance (m)": m["grid_distance_m"] if m["grid_distance_m"] == m["grid_distance_m"] else None,
            "Hatch stroke (m)": round(e["hatch"]["stroke_len"], 2) if e["family"] == "core_wall" else None,
            "Stipple dots": e.get("dots"), "Raft thickness (cm)": e.get("thickness"),
            "bbox x0 (px)": round(x0 * s, 1), "bbox y0 (px)": round(y0 * s, 1),
            "bbox x1 (px)": round(x1 * s, 1), "bbox y1 (px)": round(y1 * s, 1),
            "Matches verified detection": ("yes" if ref else "new") if e["family"] in ("column", "wall", "footing", "mass_concrete_pad") else "n/a",
            "Verified subclass": ref["subclass"] if ref else "", "Note": e.get("ref_note", ""),
            "(validation only) colour": str(tuple(e["style"][0])) if e["style"][0] else "",
            "(validation only) layer": e["style"][1] or "",
        })
    return rows


def write_workbook(res, path, dpi):
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    F = "Arial"
    head_font = Font(name=F, bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="2F4F6F")
    body = Font(name=F, size=10)
    bold = Font(name=F, size=10, bold=True)
    title = Font(name=F, size=14, bold=True)
    input_fill = PatternFill("solid", fgColor="FFFF00")
    thin = Side(style="thin", color="BBBBBB")
    wrap = Alignment(wrap_text=True, vertical="top")

    def header(ws, row, names, widths=None):
        for i, n in enumerate(names, 1):
            c = ws.cell(row=row, column=i, value=n)
            c.font, c.fill, c.alignment = head_font, head_fill, Alignment(wrap_text=True, vertical="center")
            if widths:
                ws.column_dimensions[get_column_letter(i)].width = widths[i - 1]
        ws.row_dimensions[row].height = 32

    wb = Workbook()
    rows = element_rows(res, dpi)
    N = len(rows) + 1                       # last data row on the Elements sheet
    E = lambda name: f"Elements!${col_letter(name)}$2:${col_letter(name)}${N}"

    # --- README ------------------------------------------------------------------------------
    ws = wb.active
    ws.title = "README"
    calib = res["calib"]
    lines = [
        ("Structural element catalogue, detected without colour or CAD layer", title),
        ("", body),
        ("What this is", bold),
        ("Every element found on the plan by geometry and text alone: drawing structure (solid fill, cross-hatch, closed "
         "outline, parallel line pair, dash chain), size, relative position and the plan's label grammar. Colour and layer "
         "appear only in the two '(validation only)' columns, to grade the result against the verified, colour-based "
         "detections (out/detections.json from detect_structures.py).", body),
        ("", body),
        ("Sheets", bold),
        ("Summary by building - one row per building, element family and type: count, size range, totals, consistency. Formulas over 'Elements'.", body),
        ("Elements - one row per detected element with its dimensions, colour-free marks and pixel box (150 dpi page render).", body),
        ("Marks by family - share of each family carrying each colour-free mark, and the same share among all other families.", body),
        ("P vs V - the EN 1992-1-1 column/wall test against the verified classes; threshold editable (yellow cell).", body),
        ("LG junctions - what sits where two grade beams meet, per building.", body),
        ("Validation - colour-free detections graded against the verified set, and how the band widths were chosen.", body),
        ("", body),
        ("Scale", bold),
        (f"1/{calib['scale_denominator']:.1f} ({calib['pt_per_m']:.3f} PDF points per metre), fitted from "
         f"{calib['inliers']} of {calib['samples']} dimension texts against their dimension lines, median error "
         f"{calib['median_error_cm']:.2f} cm.", body),
        ("", body),
        ("Sources for the rules", bold),
        ("Column vs wall: EN 1992-1-1 (Eurocode 2) 5.3.1(7) - a member whose section depth does not exceed 4 times its width "
         "is a column, otherwise it is treated as a wall. Applied to sizes rounded to the centimetre.", body),
        ("Paint order: AutoCAD writes the pieces of one drawn object one after another in the PDF, so consecutive filled "
         "pieces that touch are grouped into one object. This separates a column from the wall it touches without colour.", body),
        ("Band widths: the number written in LG / SF labels that has a parallel line pair of exactly that width next to most "
         "of those labels (see Validation).", body),
        ("", body),
        ("Caveats", bold),
        ("Buildings for elements without a building number in their own label are the majority of the 5 nearest labels that "
         "have one ('Building source' gives the agreement).", body),
        ("Grade beams are counted as spans between supports (footings, walls, columns, pads); strip footings as continuous runs.", body),
        ("Raft extents come from the corner diagonals drawn to the 'Radier' stamp; sump pits are sized from their text.", body),
        ("Walls drawn as one long object together with the wall they continue are split into straight stretches; a short "
         "V wall in line with a longer wall of another thickness can stay merged, and its V label then finds no wall of "
         "its size (listed on 'Validation').", body),
        ("At building joints two identical walls of neighbouring buildings stand side by side; their labels can swap.", body),
        ("Unlabelled walls standing on a raft are reported as 'core_wall_solid' (lift and stair cores carry no V mark on "
         "this plan); unlabelled walls elsewhere stay 'shear_wall' with type '(none)'.", body),
        ("Formulas were written with openpyxl and are set to recalculate when the file is opened. LibreOffice was not "
         "available to pre-compute them; they were evaluated with the Python 'formulas' engine instead (no errors) and "
         "spot-checked against pandas.", body),
    ]
    for i, (t, f) in enumerate(lines, 1):
        c = ws.cell(row=i, column=1, value=t)
        c.font, c.alignment = f, Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 130

    # --- Elements --------------------------------------------------------------------------
    ws = wb.create_sheet("Elements")
    header(ws, 1, [c for c, _ in ELEMENT_COLS], [w for _, w in ELEMENT_COLS])
    Lc, Wc, Fc = col_letter("Length (cm)"), col_letter("Width (cm)"), col_letter("Family")
    for r, row in enumerate(rows, 2):
        for ci, (name, _) in enumerate(ELEMENT_COLS, 1):
            v = row[name]
            if name == "L/W":
                v = f'=IF(AND(ISNUMBER({Lc}{r}),ISNUMBER({Wc}{r}),{Wc}{r}>0),{Lc}{r}/{Wc}{r},"")'
            elif name == "EC2 test":
                v = (f'=IF(OR({Fc}{r}="column",{Fc}{r}="wall"),IF({Lc}{r}<=\'P vs V\'!$B$3*{Wc}{r},"column","wall"),"")')
            cell = ws.cell(row=r, column=ci, value=v)
            cell.font = body
        ws.cell(row=r, column=[c for c, _ in ELEMENT_COLS].index("L/W") + 1).number_format = "0.00"
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(ELEMENT_COLS))}{N}"

    # --- Summary by building ------------------------------------------------------------------
    ws = wb.create_sheet("Summary by building", 1)
    names = ["Building", "Family", "Subclass", "Type key", "Count", "Length min (cm)", "Length max (cm)",
             "Width min (cm)", "Width max (cm)", "Same section throughout?", "Total span/run (m)", "Total area (m²)",
             "Colour-free identifying marks"]
    header(ws, 1, names, [9, 17, 19, 16, 7, 9, 9, 9, 9, 11, 10, 10, 90])
    combos = sorted({(r["Building"], r["Family"], r["Subclass"], r["Type key"]) for r in rows},
                    key=lambda t: (t[0], t[1], t[2], t[3]))
    sig = family_signatures(rows)
    crit = lambda r: (f'{E("Building")},$A{r},{E("Family")},$B{r},{E("Subclass")},$C{r},{E("Type key")},$D{r}')
    for r, (b, fam, sub, key) in enumerate(combos, 2):
        vals = [b, fam, sub, key,
                f"=COUNTIFS({crit(r)})",
                f'=IFERROR(_xlfn.MINIFS({E("Length (cm)")},{crit(r)}),"")',
                f'=IFERROR(_xlfn.MAXIFS({E("Length (cm)")},{crit(r)}),"")',
                f'=IFERROR(_xlfn.MINIFS({E("Width (cm)")},{crit(r)}),"")',
                f'=IFERROR(_xlfn.MAXIFS({E("Width (cm)")},{crit(r)}),"")',
                (f'=IF(OR(F{r}="",H{r}=""),"",IF(OR(B{r}="raft",B{r}="sump_pit",B{r}="crane_foundation",D{r}="(none)"),"n/a",'
                 f'IF(OR(B{r}="column",B{r}="footing",B{r}="mass_concrete_pad",AND(B{r}="wall",C{r}="shear_wall",D{r}<>"(none)")),'
                 f'IF(AND(G{r}-F{r}<=3,I{r}-H{r}<=3),"yes","no"),IF(I{r}-H{r}<=3,"yes (thickness)","no (thickness)"))))'),
                f"=SUMIFS({E('Span / run (m)')},{crit(r)})",
                f"=SUMIFS({E('Area (m²)')},{crit(r)})",
                sig.get((fam, sub), sig.get((fam, None), ""))]
        for ci, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=ci, value=v)
            c.font = body
            if ci == 13:
                c.alignment = wrap
        ws.cell(row=r, column=11).number_format = "0.0"
        ws.cell(row=r, column=12).number_format = "0.0"
    last = len(combos) + 1
    tot = last + 2
    ws.cell(row=tot, column=1, value="All buildings").font = bold
    ws.cell(row=tot, column=5, value=f"=SUM(E2:E{last})").font = bold
    ws.cell(row=tot, column=11, value=f"=SUM(K2:K{last})").font = bold
    ws.cell(row=tot, column=12, value=f"=SUM(L2:L{last})").font = bold
    ws.cell(row=tot + 1, column=1, value="'Same section throughout?' allows 3 cm between the smallest and largest element of the "
                                          "row. Columns, labelled shear walls, footings and pads are compared on length and width; "
                                          "walls and bands whose length varies by nature on thickness only; rafts, pits and crane "
                                          "groups, and rows of unlabelled elements (a mix of types), do not apply.").font = Font(name=F, size=9, italic=True)
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = f"A1:M{last}"

    # --- Marks by family ----------------------------------------------------------------------------
    ws = wb.create_sheet("Marks by family")
    marks = ["On raft", "Touches hatched wall", "Touches solid element", "In strip band", "On grade beam", "At LG junction"]
    header(ws, 1, ["Family", "Subclass", "Count"] + [f"{m} (share)" for m in marks] +
           [f"{m} (share in other families)" for m in marks] + ["Signature (colour-free)"],
           [17, 19, 7] + [10] * (2 * len(marks)) + [90])
    fams = sorted({(r["Family"], r["Subclass"]) for r in rows})
    for r, (fam, sub) in enumerate(fams, 2):
        ws.cell(row=r, column=1, value=fam).font = body
        ws.cell(row=r, column=2, value=sub).font = body
        ws.cell(row=r, column=3, value=f'=COUNTIFS({E("Family")},$A{r},{E("Subclass")},$B{r})').font = body
        for i, mk in enumerate(marks):
            c = ws.cell(row=r, column=4 + i,
                        value=f'=IFERROR(COUNTIFS({E("Family")},$A{r},{E("Subclass")},$B{r},{E(mk)},"yes")/$C{r},0)')
            c.number_format, c.font = "0%", body
            c2 = ws.cell(row=r, column=4 + len(marks) + i,
                         value=f'=IFERROR((COUNTIFS({E(mk)},"yes")-COUNTIFS({E("Family")},$A{r},{E("Subclass")},$B{r},{E(mk)},"yes"))'
                               f'/(COUNTA({E("ID")})-$C{r}),0)')
            c2.number_format, c2.font = "0%", body
        c = ws.cell(row=r, column=4 + 2 * len(marks), value=sig.get((fam, sub), sig.get((fam, None), "")))
        c.font, c.alignment = body, wrap
    ws.freeze_panes = "D2"

    # --- P vs V -----------------------------------------------------------------------------------
    ws = wb.create_sheet("P vs V")
    ws["A1"] = "Column (P) vs wall (V) without colour: EN 1992-1-1 5.3.1(7) section test against the verified classes"
    ws["A1"].font = title
    ws["A3"], ws["B3"] = "Column if length <= this x width", 4
    ws["A3"].font, ws["B3"].font = bold, Font(name=F, size=10, color="0000FF")
    ws["B3"].fill = input_fill
    ws["B3"].comment = Comment("EN 1992-1-1 5.3.1(7): a column's section depth does not exceed 4 times its width. "
                               "Change to test other thresholds; the EC2 test column on 'Elements' follows.", "element_stats")
    ws["A5"] = "Verified class (colour pipeline) vs EC2 test (colour-free)"
    ws["A5"].font = bold
    for i, h in enumerate(["Verified subclass", "EC2 says column", "EC2 says wall", "Total"], 1):
        c = ws.cell(row=6, column=i, value=h)
        c.font, c.fill = head_font, head_fill
    for r, sub in enumerate(["column", "shear_wall", "retaining_wall"], 7):
        ws.cell(row=r, column=1, value=sub).font = body
        ws.cell(row=r, column=2, value=f'=COUNTIFS({E("Verified subclass")},$A{r},{E("EC2 test")},"column")').font = body
        ws.cell(row=r, column=3, value=f'=COUNTIFS({E("Verified subclass")},$A{r},{E("EC2 test")},"wall")').font = body
        ws.cell(row=r, column=4, value=f"=B{r}+C{r}").font = body
    ws["A11"] = "Agreement"
    ws["A11"].font = bold
    ws["B11"] = "=(B7+C8+C9)/SUM(D7:D9)"
    ws["B11"].number_format = "0.0%"
    ws["A13"] = "Section ratio L/W by verified class"
    ws["A13"].font = bold
    for i, h in enumerate(["Verified subclass", "Min L/W", "Max L/W", "Min length (cm)", "Max length (cm)"], 1):
        c = ws.cell(row=14, column=i, value=h)
        c.font, c.fill = head_font, head_fill
    for r, sub in enumerate(["column", "shear_wall", "retaining_wall"], 15):
        ws.cell(row=r, column=1, value=sub).font = body
        for ci, (fn, colname) in enumerate((("MINIFS", "L/W"), ("MAXIFS", "L/W"), ("MINIFS", "Length (cm)"),
                                            ("MAXIFS", "Length (cm)")), 2):
            c = ws.cell(row=r, column=ci, value=f'=IFERROR(_xlfn.{fn}({E(colname)},{E("Verified subclass")},$A{r}),"")')
            c.number_format, c.font = ("0.00" if colname == "L/W" else "0"), body
    ws["A19"] = "Gap between the stubbiest wall and the most elongated column"
    ws["A19"].font = bold
    ws["B19"] = "=B16-C15"
    ws["B19"].number_format = "0.00"
    ws["A21"] = "Colour-free families of all filled rectangles (including walls the colour pipeline missed or left unsplit)"
    ws["A21"].font = bold
    for i, h in enumerate(["Family", "Count", "of which verified", "new (not in verified set)"], 1):
        c = ws.cell(row=22, column=i, value=h)
        c.font, c.fill = head_font, head_fill
    for r, fam in enumerate(["column", "wall"], 23):
        ws.cell(row=r, column=1, value=fam).font = body
        ws.cell(row=r, column=2, value=f'=COUNTIFS({E("Family")},$A{r})').font = body
        ws.cell(row=r, column=3, value=f'=COUNTIFS({E("Family")},$A{r},{E("Matches verified detection")},"yes")').font = body
        ws.cell(row=r, column=4, value=f'=COUNTIFS({E("Family")},$A{r},{E("Matches verified detection")},"new")').font = body
    ws.column_dimensions["A"].width = 58
    for c in "BCDE":
        ws.column_dimensions[c].width = 16

    # --- LG junctions -------------------------------------------------------------------------------
    ws = wb.create_sheet("LG junctions")
    header(ws, 1, ["Junction", "Building", "What sits there", "Everything within 60 cm", "Wall type", "x (px)", "y (px)",
                   "Nearest element", "Distance to it (m)"], [9, 9, 22, 30, 12, 9, 9, 16, 10])
    s = dpi / 72
    jr = sorted(res["junctions"], key=lambda j: (str(j["building"]), j["what"]))
    for r, j in enumerate(jr, 2):
        for ci, v in enumerate([r - 1, j["building"] or "unknown", j["what"], j["all"], j["wall_key"],
                                round(j["x"] * s, 1), round(j["y"] * s, 1), j["nearest"], j["nearest_m"]], 1):
            ws.cell(row=r, column=ci, value=v).font = body
    jl = len(jr) + 1
    cats = ["wall", "footing carrying wall", "column", "footing carrying column", "core_wall", "mass_concrete_pad",
            "footing carrying nothing", "raft", "nothing"]   # core_wall = hatched core wall
    blds = sorted({j["building"] or "unknown" for j in jr})
    top, C0 = 1, 11
    ws.cell(row=top, column=C0, value="Building").font = head_font
    ws.cell(row=top, column=C0).fill = head_fill
    for ci, cat in enumerate(cats + ["Total"], C0 + 1):
        c = ws.cell(row=top, column=ci, value=cat)
        c.font, c.fill = head_font, head_fill
        ws.column_dimensions[get_column_letter(ci)].width = 11
    bl = get_column_letter(C0)
    for ri, b in enumerate(blds + ["All"], top + 1):
        ws.cell(row=ri, column=C0, value=b).font = bold if b == "All" else body
        for ci, cat in enumerate(cats, C0 + 1):
            f = (f'=COUNTIFS($B$2:$B${jl},${bl}{ri},$C$2:$C${jl},{get_column_letter(ci)}${top})' if b != "All"
                 else f'=COUNTIFS($C$2:$C${jl},{get_column_letter(ci)}${top})')
            ws.cell(row=ri, column=ci, value=f).font = body
        ws.cell(row=ri, column=C0 + 1 + len(cats),
                value=f"=SUM({get_column_letter(C0 + 1)}{ri}:{get_column_letter(C0 + len(cats))}{ri})").font = bold
    ar = top + len(blds) + 1
    ws.cell(row=ar + 1, column=C0, value="Share").font = bold
    for ci in range(C0 + 1, C0 + 1 + len(cats)):
        c = ws.cell(row=ar + 1, column=ci,
                    value=f"=IFERROR({get_column_letter(ci)}{ar}/{get_column_letter(C0 + 1 + len(cats))}{ar},0)")
        c.number_format, c.font = "0%", body
    ws.cell(row=ar + 3, column=C0, value="Junctions with no element within 60 cm, by distance to the nearest one").font = bold
    for i, (lo, hi) in enumerate(((0.6, 1.0), (1.0, 2.0), (2.0, 99.0))):
        ws.cell(row=ar + 4 + i, column=C0, value=f"{lo}-{hi if hi < 99 else '+'} m").font = body
        ws.cell(row=ar + 4 + i, column=C0 + 1,
                value=f'=COUNTIFS($I$2:$I${jl},">="&{lo},$I$2:$I${jl},"<"&{hi})').font = body
    ws.column_dimensions[bl].width = 10
    ws.freeze_panes = "A2"

    # --- Validation --------------------------------------------------------------------------------
    ws = wb.create_sheet("Validation")
    ws["A1"] = "Colour-free detection graded against the verified (colour-based) detections"
    ws["A1"].font = title
    hdr = ["Family", "Verified elements", "Found colour-free", "Recall", "Found but as another family",
           "Colour-free total", "Not in verified set"]
    for i, h in enumerate(hdr, 1):
        c = ws.cell(row=3, column=i, value=h)
        c.font, c.fill, c.alignment = head_font, head_fill, Alignment(wrap_text=True)
    for r, g in enumerate(res["grading"], 4):
        vals = [g["family"], g["verified"], g["found_colour_free"], f"=IFERROR(C{r}/B{r},0)", g["found_as_other_family"],
                g["colour_free_total"], g["not_in_verified_set"]]
        for ci, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=ci, value=v)
            c.font = body
        ws.cell(row=r, column=4).number_format = "0.0%"
    r0 = 4 + len(res["grading"]) + 1
    ws.cell(row=r0, column=1, value="Numbers above are program output (counts), not spreadsheet calculations; "
                                   "'Not in verified set' lists elements the colour pipeline missed or did not split "
                                   "(see the Note column on Elements).").font = Font(name=F, size=9, italic=True)
    r0 += 2
    ws.cell(row=r0, column=1, value="How band widths were read from the labels").font = bold
    for i, h in enumerate(["Label type", "Labels", "Width chosen (m)", "Labels with such a band within 1.5 m", "Share"], 1):
        c = ws.cell(row=r0 + 1, column=i, value=h)
        c.font, c.fill = head_font, head_fill
    for r, (key, b) in enumerate(sorted(res["log"].get("bands", {}).items()), r0 + 2):
        for ci, v in enumerate([key, b["labels"], b["width_m"], b["labels_with_band"], f"=IFERROR(D{r}/B{r},0)"], 1):
            ws.cell(row=r, column=ci, value=v).font = body
        ws.cell(row=r, column=5).number_format = "0%"
    rb = r0 + 3 + len(res["log"].get("bands", {}))
    ws.cell(row=rb, column=1, value="Band lines by CAD layer (validation only: the detector never reads layers)").font = bold
    for i, (kind, v) in enumerate(sorted(res["log"].get("band_layers", {}).items()), rb + 1):
        ws.cell(row=i, column=1, value=f"{kind}: {v['runs']} runs, {v['total_run_m']} m, {v['lines']} lines").font = body
        ws.cell(row=i, column=2, value=", ".join(f"{l or '(none)'} {n}" for l, n in v["top_layers"])).font = body
    r1 = rb + 2 + len(res["log"].get("band_layers", {}))
    ws.cell(row=r1, column=1, value="Other counts").font = bold
    other = [("Filled rectangular objects (paint-order groups)", res["log"]["filled"]["rectangular_objects"]),
             ("Filled objects drawn twice (removed)", res["log"]["filled"]["drawn_twice"]),
             ("Wall stretches cut from thin bent fills", res["log"]["filled"].get("stretches_from_bent_walls", 0)),
             ("Cross-hatched regions (walls)", res["log"]["hatched_regions"]["cross"]),
             ("Other hatch patterns (stairs, ramps; not walls)", res["log"]["hatched_regions"]["other_patterns"]),
             ("Grid axes found (dash chain + named bubble)", len(res["grid"])),
             ("VP labels next to a retaining wall of the stated thickness", res["log"]["vp_labels"].get("thickness agrees", 0)),
             ("VP labels next to a retaining wall of another thickness", res["log"]["vp_labels"].get("thickness differs", 0)),
             ("VP labels with no retaining wall within 1.5 m", res["log"]["vp_labels"].get("no retaining wall within 1.5 m", 0)),
             ("Found outside the drawing area (no labelled element within 25 m; title block, legend): " +
              ", ".join(f"{n} {f}" for f, n in res["log"].get("outside_plan", {}).items()),
              sum(res["log"].get("outside_plan", {}).values())),
             ("Dashed-line pieces ignored as element edges", res["log"]["dashed_pieces"]),
             ("Axis pen width learned from the dashed lines (pt); lines in this pen are not element edges",
              res["log"]["axis_pen_pt"]),
             ("Labels repeating a type already on the element next to them", len(res["log"]["labels_repeating"])),
             ("Labels with no element of their type nearby: " + ", ".join(res["log"]["labels_without_shape"]),
              len(res["log"]["labels_without_shape"]))]
    for r, (t, v) in enumerate(other, r1 + 1):
        ws.cell(row=r, column=1, value=t).font = body
        ws.cell(row=r, column=2, value=v).font = body
    ws.column_dimensions["A"].width = 52
    for c in "BCDEFG":
        ws.column_dimensions[c].width = 16

    wb.calculation.fullCalcOnLoad = True
    wb.save(path)
    return rows


def family_signatures(rows):
    """Plain-language colour-free signature of each family, computed from the elements."""
    out = {}
    by = collections.defaultdict(list)
    for r in rows:
        by[(r["Family"], r["Subclass"])].append(r)
    pct = lambda rs, col: 100 * sum(1 for r in rs if r[col] == "yes") / len(rs)
    for (fam, sub), rs in by.items():
        Ls = [r["Length (cm)"] for r in rs if r["Length (cm)"]]
        Ws = [r["Width (cm)"] for r in rs if r["Width (cm)"]]
        size = f"{min(Ls)}-{max(Ls)} x {min(Ws)}-{max(Ws)} cm" if Ls and Ws else ""
        ratio = ""
        if Ls and Ws and fam in ("column", "wall"):
            q = [r["Length (cm)"] / r["Width (cm)"] for r in rs if r["Length (cm)"] and r["Width (cm)"]]
            ratio = f"; L/W {min(q):.2f}-{max(q):.2f}"
        parts = [collections.Counter(r["Drawing type"] for r in rs).most_common(1)[0][0], size + ratio]
        if sub == "core_wall_hatched":
            st = sorted({r["Hatch stroke (m)"] for r in rs if r["Hatch stroke (m)"]})
            parts.append(f"hatch strokes {st} m at +/-45 deg; thickness = local spread of stroke ends")
            parts.append(f"on a raft {pct(rs, 'On raft'):.0f}%; no V label")
        elif fam in ("column", "wall"):
            inf = sum(1 for r in rs if r["Inside footing"]) * 100 / len(rs)
            parts.append(f"inside a footing outline {inf:.0f}%")
            parts.append(f"on a raft {pct(rs, 'On raft'):.0f}%")
            parts.append(f"touching a hatched core wall {pct(rs, 'Touches hatched wall'):.0f}%")
            parts.append(f"in a strip-footing band {pct(rs, 'In strip band'):.0f}%")
            parts.append(f"at a grade-beam junction {pct(rs, 'At LG junction'):.0f}%")
        if fam == "footing":
            sup = collections.Counter(r["Supports inside"] for r in rs).most_common(3)
            parts.append("carries " + ", ".join(f"{s or 'nothing'} ({n})" for s, n in sup))
            parts.append(f"grade beam crossing it {pct(rs, 'On grade beam'):.0f}%")
        if fam == "mass_concrete_pad":
            d = [r["Stipple dots"] for r in rs]
            parts.append(f"{min(d)}-{max(d)} stipple dots inside; on a grade beam {pct(rs, 'On grade beam'):.0f}%; "
                         f"at a grade-beam junction {pct(rs, 'At LG junction'):.0f}%")
        if fam in ("grade_beam", "strip_footing"):
            parts.append("width stated in its label")
        if fam == "raft":
            parts.append("circle stamp 'Radier' + thickness; diagonals from its corners to the stamp")
        if fam == "sump_pit":
            parts.append("'FOSSE DE RELEVAGE' text with LxWxH, 'RADIER 25' stamp")
        out[(fam, sub)] = "; ".join(p for p in parts if p)
    return out


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("pdf")
    ap.add_argument("--config", default=str(Path(__file__).with_name("plan_config.toml")))
    ap.add_argument("--reference", default="out/detections.json", help="verified detections used only for grading")
    ap.add_argument("--out", default="out")
    ap.add_argument("--dpi", type=int, default=150)
    args = ap.parse_args()
    cfg = tomllib.loads(Path(args.config).read_text())
    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)
    res = build(args.pdf, cfg, args.reference)
    rows = write_workbook(res, out / "element_catalogue.xlsx", args.dpi)
    print(f"{len(rows)} elements -> {out / 'element_catalogue.xlsx'}")
    print("families:", dict(collections.Counter(r["Family"] for r in rows)))
    for g in res["grading"]:
        print("grading:", g)
    print("junctions:", dict(collections.Counter(j["what"] for j in res["junctions"])))
    return 0


if __name__ == "__main__":
    sys.exit(main())
